#!/usr/bin/env python3
"""Prepare and screen reported numeric data for research-integrity signals.

Standard library only (openpyxl is needed just for .xlsx input).  Like
audit_figures.py, it produces reviewable candidates, never verdicts.
"""

from __future__ import annotations

import argparse
import csv
import io
import itertools
import json
import math
import re
import statistics
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any


TABLE_SUFFIXES = {".csv", ".tsv", ".txt", ".xlsx", ".xlsm"}
NUMBER_RE = re.compile(r"[+-]?(?:\d+\.?\d*|\.\d+)(?:[eE][+-]?\d+)?")
BENFORD = [math.log10(1 + 1 / digit) for digit in range(1, 10)]
STAT_RE = re.compile(
    r"""(?<![A-Za-z])(?P<test>t|F|r|z|chi2|chi-square|χ2|χ²|X2)\s*
    (?:\(\s*(?P<df1>\d+(?:\.\d+)?)\s*(?:,\s*(?P<df2>\d+(?:\.\d+)?))?\s*(?:,\s*N\s*=\s*\d+\s*)?\))?
    \s*=\s*(?P<stat>-?(?:\d+\.?\d*|\.\d+))\s*[,;]\s*p\s*(?P<cmp><=|>=|[<>=≤≥])\s*(?P<p>\d*\.?\d+)""",
    re.IGNORECASE | re.VERBOSE,
)
MIN_DIGIT_VALUES = 50  # expected >= 5 per terminal digit for the chi-square approximation
MIN_BENFORD_VALUES = 100


def write_json(path: Path, value: Any) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def require_new_workspace(path: Path) -> None:
    if path.exists() and any(path.iterdir()):
        raise SystemExit(f"refusing to overwrite non-empty audit workspace: {path}")
    path.mkdir(parents=True, exist_ok=True)


def slug(value: str, fallback: str = "table") -> str:
    cleaned = re.sub(r"[^A-Za-z0-9_.-]+", "-", value).strip("-.")
    return cleaned or fallback


def column_letter(index: int) -> str:
    letters = ""
    index += 1
    while index:
        index, rest = divmod(index - 1, 26)
        letters = chr(65 + rest) + letters
    return letters


# --- number text ---------------------------------------------------------


def clean_number(text: Any) -> str | None:
    """Return the reported number text, or None when the cell is not a plain number."""
    value = str(text).strip().replace("−", "-")  # PDF tables often use a Unicode minus
    return value if NUMBER_RE.fullmatch(value) else None


def decimals(text: str) -> int:
    mantissa = re.split("[eE]", text)[0]
    return len(mantissa.split(".", 1)[1]) if "." in mantissa else 0


def digit_string(text: str, places: int) -> str | None:
    """Digits of `text` padded to `places` decimals, without sign, point, or leading zeros.

    Padding restores trailing zeros that spreadsheets and exporters drop, so a
    column reported at two decimals keeps its terminal zeros.  Scientific
    notation is excluded from digit-level tests.
    """
    if "e" in text.lower():
        return None
    whole, _, fraction = text.lstrip("+-").partition(".")
    return (whole + fraction.ljust(places, "0")).lstrip("0")


def series_precision(values: list[str]) -> int:
    return max((decimals(value) for value in values), default=0)


# --- statistics ----------------------------------------------------------


def chi2_sf(x: float, df: int) -> float:
    """Upper tail of the chi-square distribution for integer df (closed form)."""
    if x <= 0:
        return 1.0
    half = x / 2
    total = math.exp(-half) if df % 2 == 0 else math.erfc(math.sqrt(half))
    for m in range(2 - df % 2, df, 2):
        total += math.exp(m / 2 * math.log(half) - half - math.lgamma(m / 2 + 1))
    return min(1.0, total)


def chi_square(observed: list[int], probabilities: list[float]) -> tuple[float, float]:
    n = sum(observed)
    statistic = sum((o - n * p) ** 2 / (n * p) for o, p in zip(observed, probabilities))
    return statistic, chi2_sf(statistic, len(observed) - 1)


def betainc(a: float, b: float, x: float) -> float:
    """Regularized incomplete beta I_x(a, b) by Lentz's continued fraction."""
    if x <= 0:
        return 0.0
    if x >= 1:
        return 1.0
    if x > (a + 1) / (a + b + 2):
        return 1.0 - betainc(b, a, 1.0 - x)
    front = math.exp(a * math.log(x) + b * math.log1p(-x) - math.lgamma(a) - math.lgamma(b) + math.lgamma(a + b)) / a
    f, c, d = 1.0, 1.0, 0.0
    for i in range(400):
        m = i // 2
        if i == 0:
            numerator = 1.0
        elif i % 2 == 0:
            numerator = m * (b - m) * x / ((a + 2 * m - 1) * (a + 2 * m))
        else:
            numerator = -(a + m) * (a + b + m) * x / ((a + 2 * m) * (a + 2 * m + 1))
        d = 1.0 + numerator * d
        d = 1.0 / (d if abs(d) > 1e-300 else 1e-300)
        c = 1.0 + numerator / c
        c = c if abs(c) > 1e-300 else 1e-300
        f *= c * d
        if abs(1.0 - c * d) < 1e-14:
            break
    return front * (f - 1.0)


def p_value(test: str, statistic: float, df1: float | None, df2: float | None) -> float:
    """Two-sided p for t, r, and z; upper-tail p for F and chi-square."""
    statistic = abs(statistic)
    if test == "t":
        return betainc(df1 / 2, 0.5, df1 / (df1 + statistic**2))
    if test == "r":
        return 0.0 if statistic >= 1 else p_value("t", statistic * math.sqrt(df1 / (1 - statistic**2)), df1, None)
    if test == "z":
        return math.erfc(statistic / math.sqrt(2))
    if test == "F":
        return betainc(df2 / 2, df1 / 2, df2 / (df2 + df1 * statistic))
    return chi2_sf(statistic, int(df1))


def binom_sf(k: int, n: int, p: float) -> float:
    """P(X >= k) for X ~ Binomial(n, p)."""
    return min(1.0, sum(math.comb(n, j) * p**j * (1 - p) ** (n - j) for j in range(max(k, 0), n + 1)))


def benjamini_hochberg(pvalues: list[float]) -> list[float]:
    order = sorted(range(len(pvalues)), key=pvalues.__getitem__)
    qvalues = [1.0] * len(pvalues)
    running = 1.0
    for rank in range(len(order), 0, -1):
        index = order[rank - 1]
        running = min(running, pvalues[index] * len(pvalues) / rank)
        qvalues[index] = running
    return qvalues


# --- checks --------------------------------------------------------------


def last_digits(values: list[str]) -> list[str]:
    """Last digits of one series' values with >= 3 significant digits, padded to its precision."""
    places = series_precision(values)
    return [d[-1] for d in (digit_string(v, places) for v in values) if d and len(d) >= 3]


def spans_unit(values: list[str]) -> bool:
    """Decimal parts are only near-uniform when the values span at least one whole unit."""
    numbers = [float(v) for v in values]
    return max(numbers) - min(numbers) >= 1


def terminal_digit(digits: list[str]) -> dict[str, Any] | None:
    """Chi-square uniformity of terminal digits (pad per series before pooling)."""
    if len(digits) < MIN_DIGIT_VALUES:
        return None
    counts = Counter(digits)
    observed = [counts.get(str(digit), 0) for digit in range(10)]
    statistic, p = chi_square(observed, [0.1] * 10)
    top = max(range(10), key=observed.__getitem__)
    return {
        "n": len(digits),
        "statistic": round(statistic, 3),
        "p": p,
        "detail": f"last-digit counts 0-9 {observed}; most common {top} "
        f"({observed[top] / len(digits):.0%}, expected 10%)",
    }


def benford(values: list[str]) -> dict[str, Any] | None:
    """First-digit Benford conformity; only for positive data spanning >= 2 orders of magnitude."""
    numbers = [abs(float(v)) for v in values if float(v) != 0]
    if len(numbers) < MIN_BENFORD_VALUES or max(numbers) / min(numbers) < 100:
        return None
    first = Counter(int(f"{number:e}"[0]) for number in numbers)
    observed = [first.get(digit, 0) for digit in range(1, 10)]
    statistic, p = chi_square(observed, BENFORD)
    mad = sum(abs(o / len(numbers) - e) for o, e in zip(observed, BENFORD)) / 9
    conformity = "close" if mad < 0.006 else "acceptable" if mad < 0.012 else "marginal" if mad < 0.015 else "nonconforming"
    return {
        "n": len(numbers),
        "statistic": round(statistic, 3),
        "p": p,
        "mad": mad,
        "detail": f"first-digit counts 1-9 {observed}; MAD {mad:.4f} ({conformity}, Nigrini thresholds)",
    }


def decimal_repetition(values: list[str]) -> dict[str, Any] | None:
    """Most repeated decimal part within one series, against uniform decimal parts."""
    places = series_precision(values)
    if places < 2 or not spans_unit(values):
        return None
    parts = [d[-places:].rjust(places, "0") for d in (digit_string(v, places) for v in values) if d is not None]
    if len(parts) < 5:
        return None
    part, count = Counter(parts).most_common(1)[0]
    choices = 10**places
    return {
        "n": len(parts),
        "statistic": count,
        "p": min(1.0, choices * binom_sf(count, len(parts), 1 / choices)),
        "detail": f"decimal part .{part} appears {count}/{len(parts)} times",
    }


def fixed_relation(first: list[str], second: list[str]) -> dict[str, Any] | None:
    """Constant difference, ratio, or exact linear map between paired rows, within rounding."""
    a = [float(v) for v in first]
    b = [float(v) for v in second]
    half_a = 0.5 * 10 ** -series_precision(first)
    half_b = 0.5 * 10 ** -series_precision(second)
    if a == b:
        return {"relation": "identical", "detail": "all paired values identical"}
    differences = [y - x for x, y in zip(a, b)]
    if max(differences) - min(differences) <= 2 * (half_a + half_b) + 1e-9:
        exact = max(differences) - min(differences) < 1e-9
        return {
            "relation": "difference",
            "detail": f"B - A = {statistics.fmean(differences):.6g} for every row"
            + (" (exact)" if exact else " (within rounding)"),
        }
    if all(abs(x) > half_a for x in a):
        low, high = -math.inf, math.inf
        for x, y in zip(a, b):
            corners = [(y + sy) / (x + sx) for sx in (-half_a, half_a) for sy in (-half_b, half_b)]
            low, high = max(low, min(corners)), min(high, max(corners))
        if low <= high + 1e-12:
            ratios = [y / x for x, y in zip(a, b)]
            exact = max(ratios) - min(ratios) < 1e-9 * max(1.0, abs(ratios[0]))
            return {
                "relation": "ratio",
                "detail": f"B / A = {statistics.fmean(ratios):.6g} for every row"
                + (" (exact)" if exact else " (within rounding)"),
            }
    if len(a) >= 5:
        slope, intercept = statistics.linear_regression(a, b)
        residual = max(abs(y - (slope * x + intercept)) for x, y in zip(a, b))
        if residual <= half_b + abs(slope) * half_a + 1e-9:
            return {
                "relation": "linear",
                "detail": f"B = {slope:.6g} * A + {intercept:.6g}; max residual {residual:.3g} is within rounding",
            }
    return None


def decimal_match(first: list[str], second: list[str]) -> dict[str, Any] | None:
    """Rows whose decimal parts coincide across two series."""
    places = min(series_precision(first), series_precision(second))
    if places < 1 or not spans_unit(first) or not spans_unit(second):
        return None
    pairs = [(digit_string(x, places), digit_string(y, places)) for x, y in zip(first, second)]
    pairs = [(x, y) for x, y in pairs if x is not None and y is not None]
    matches = sum(x[-places:].rjust(places, "0") == y[-places:].rjust(places, "0") for x, y in pairs)
    return {
        "n": len(pairs),
        "statistic": matches,
        "p": binom_sf(matches, len(pairs), 10**-places),
        "detail": f"{matches}/{len(pairs)} rows share the same {places}-digit decimal part",
    }


def shared_runs(series: list[dict], min_run: int) -> list[dict[str, Any]]:
    """Runs of >= min_run identical consecutive values repeated across or within series."""
    values = {item["id"]: [float(v) for v in item["values"]] for item in series}
    index: dict[tuple, list[tuple[str, int]]] = defaultdict(list)
    for series_id, numbers in values.items():
        for start in range(len(numbers) - min_run + 1):
            window = tuple(numbers[start : start + min_run])
            if len(set(window)) >= 3:  # ponytail: skips runs of repeated constants like 0,0,0,1
                index[window].append((series_id, start))
    seen, runs = set(), []
    for places in index.values():
        for (first, i), (second, j) in itertools.combinations(places, 2):
            x, y = values[first], values[second]
            while i and j and x[i - 1] == y[j - 1]:
                i, j = i - 1, j - 1
            if (first, i, second, j) in seen:
                continue
            seen.add((first, i, second, j))
            length = 0
            while i + length < len(x) and j + length < len(y) and x[i + length] == y[j + length]:
                length += 1
            if first == second and abs(i - j) < length:
                continue  # overlapping self-match from a periodic pattern
            runs.append(
                {
                    "series_a": first,
                    "series_b": second,
                    "n": length,
                    "detail": f"items {i + 1}-{i + length} of A equal items {j + 1}-{j + length} of B: "
                    + ", ".join(f"{v:g}" for v in x[i : i + length]),
                }
            )
    return runs


def grim(mean_text: str, n: int, items: int = 1, percent: bool = False) -> dict[str, Any]:
    """GRIM: can a mean of n integer responses (x items) round to the reported value?"""
    text = clean_number(mean_text)
    if text is None or n <= 0 or items <= 0:
        raise ValueError(f"invalid GRIM input: mean={mean_text!r}, n={n}, items={items}")
    places = decimals(text)
    scale = 100 if percent else 1
    total = n * items
    if total >= 10**places * scale:
        return {"testable": False, "detail": f"n x items = {total} is too large for {places} reported decimals"}
    mean = float(text) / scale
    unit = 10**-places / scale
    low = math.ceil((mean - unit / 2) * total - 1e-9)
    high = math.floor((mean + unit / 2) * total + 1e-9)
    if low <= high:
        return {"testable": True, "consistent": True, "detail": f"sum {low} / {total} is consistent"}
    nearest = [f"{k / total * scale:.{places}f}" for k in (high, low)]
    return {
        "testable": True,
        "consistent": False,
        "detail": f"no integer sum over {total} responses rounds to {text}; nearest possible {nearest[0]} or {nearest[1]}",
    }


def grimmer(mean_text: str, sd_text: str, n: int, items: int = 1) -> dict[str, Any]:
    """GRIMMER: can n integer responses (x items) have this mean and sample SD?

    Sum of squares Q = (n - 1) SD^2 + T^2 / n must be an integer with the same
    parity as the sum T, because x^2 and x are always both even or both odd.
    """
    mean, sd = clean_number(mean_text), clean_number(sd_text)
    if mean is None or sd is None or n < 2 or items <= 0:
        raise ValueError(f"invalid GRIMMER input: mean={mean_text!r}, sd={sd_text!r}, n={n}, items={items}")
    # Per-participant item sums are integers; work in those units.
    center, half_mean = float(mean) * items, 0.5 * 10 ** -decimals(mean) * items
    spread, half_sd = float(sd) * items, 0.5 * 10 ** -decimals(sd) * items
    low_q, high_q = (n - 1) * max(spread - half_sd, 0) ** 2, (n - 1) * (spread + half_sd) ** 2
    if high_q - low_q >= 2:
        return {"testable": False, "detail": f"SD {sd} is too coarse to constrain n = {n}"}
    for total in range(math.ceil((center - half_mean) * n - 1e-9), math.floor((center + half_mean) * n + 1e-9) + 1):
        squares = math.ceil(low_q + total**2 / n - 1e-9)
        squares += (squares - total) % 2
        if squares <= high_q + total**2 / n + 1e-9:
            return {"testable": True, "consistent": True, "detail": f"sum {total}, sum of squares {squares} fit"}
    return {
        "testable": True,
        "consistent": False,
        "detail": f"no integer data of n = {n} has mean {mean} and SD {sd} (sum-of-squares or parity mismatch)",
    }


def statcheck(report: str, tails: int = 2, alpha: float = 0.05) -> dict[str, Any]:
    """Recompute p from a reported test statistic and degrees of freedom, allowing for rounding."""
    match = STAT_RE.search(report.replace("\u2212", "-"))
    if not match:
        raise ValueError(f"cannot parse test report (expected e.g. 't(18) = 2.31, p = .032'): {report!r}")
    test = {"f": "F", "t": "t", "r": "r", "z": "z"}.get(match["test"].lower(), "chi2")
    df1 = float(match["df1"]) if match["df1"] else None
    df2 = float(match["df2"]) if match["df2"] else None
    if (test != "z" and df1 is None) or (test == "F") != (df2 is not None) or (test == "chi2" and not df1.is_integer()):
        raise ValueError(f"missing or invalid degrees of freedom for {test}: {report!r}")
    statistic = abs(float(match["stat"]))
    if test == "r" and statistic >= 1:
        raise ValueError(f"correlation must be below 1: {report!r}")
    half = 0.5 * 10 ** -decimals(match["stat"])
    largest = min(statistic + half, 1 - 1e-12) if test == "r" else statistic + half
    factor = 0.5 if tails == 1 and test in {"t", "r", "z"} else 1.0
    low, high = (factor * p_value(test, value, df1, df2) for value in (largest, max(statistic - half, 0.0)))
    reported, comparison = float(match["p"]), {"≤": "<", "<=": "<", "≥": ">", ">=": ">"}.get(match["cmp"], match["cmp"])
    half_p = 0.5 * 10 ** -decimals(match["p"])

    def fits(low: float, high: float) -> bool:
        if comparison == "=":
            return low <= reported + half_p and high >= reported - half_p
        return low < reported if comparison == "<" else high > reported

    consistent = fits(low, high)
    reported_significant = reported <= alpha if comparison == "<" else comparison == "=" and reported < alpha
    reported_nonsignificant = comparison != "<" and reported >= alpha
    decision_error = (reported_significant and low >= alpha) or (reported_nonsignificant and high < alpha)
    detail = f"recomputed {'one' if factor < 1 else 'two' if test in {'t', 'r', 'z'} else 'upper'}-tailed p = {factor * p_value(test, statistic, df1, df2):.4g} (rounding range {low:.3g}-{high:.3g}); reported p {comparison} {match['p']}"
    if decision_error:
        detail += f"; significance at {alpha} changes"
    if not consistent and factor == 1 and test in {"t", "r", "z"} and fits(low / 2, high / 2):
        detail += "; consistent if one-tailed"
    return {"consistent": consistent, "decision_error": decision_error, "detail": detail}


# --- prepare -------------------------------------------------------------


def xlsx_cell_text(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, float):
        fmt = cell.number_format or ""
        fixed = re.search(r"0\.(0+)", fmt)
        if fixed and "%" not in fmt and "E" not in fmt.upper():
            return f"{value:.{len(fixed.group(1))}f}"  # the precision the authors displayed
        return f"{value:.15g}"
    return str(value).strip()


def read_tables(path: Path) -> list[tuple[str, list[list[str]], set[tuple[int, int]]]]:
    """Return (table name, cell grid, formula cell coordinates) for every table in a file."""
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        try:
            import openpyxl
        except ImportError as error:
            raise SystemExit("openpyxl is required for Excel input; install it or export each sheet to CSV") from error
        values = openpyxl.load_workbook(path, read_only=True, data_only=True)
        formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
        tables = []
        for sheet in values.worksheets:
            grid = [[xlsx_cell_text(cell) for cell in row] for row in sheet.iter_rows()]
            formula_cells = {
                (r, c)
                for r, row in enumerate(formulas[sheet.title].iter_rows(values_only=True))
                for c, value in enumerate(row)
                if isinstance(value, str) and value.startswith("=")
            }
            tables.append((sheet.title, grid, formula_cells))
        values.close()
        formulas.close()
        return tables
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    if path.suffix.lower() == ".tsv":
        delimiter = "\t"
    else:
        try:
            delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t").delimiter
        except csv.Error:
            delimiter = ","
    return [(path.stem, list(csv.reader(io.StringIO(text), delimiter=delimiter)), set())]


def propose_series(source: str, table: str, grid: list[list[str]], formula_cells: set) -> list[dict]:
    """One proposal per vertical block of >= 3 consecutive numeric cells."""
    def cell(r: int, c: int) -> str:
        return grid[r][c] if c < len(grid[r]) else ""

    proposals = []
    for c in range(max((len(row) for row in grid), default=0)):
        r = 0
        while r < len(grid):
            if clean_number(cell(r, c)) is None:
                r += 1
                continue
            start = r
            while r < len(grid) and clean_number(cell(r, c)) is not None:
                r += 1
            if r - start < 3:
                continue
            values = [clean_number(cell(i, c)) for i in range(start, r)]
            above = next((cell(i, c).strip() for i in range(start - 1, -1, -1) if cell(i, c).strip()), "")
            letter = column_letter(c)
            is_index = all(float(v) == start_value for v, start_value in zip(values, itertools.count(float(values[0]))))
            proposals.append(
                {
                    "id": f"{source}:{table}!{letter}{start + 1}:{letter}{r}",
                    "label": above if clean_number(above) is None else "",
                    "source": source,
                    "values": values,
                    "formula": any((i, c) in formula_cells for i in range(start, r)),
                    "derivation_group": None,
                    "include": not is_index,
                    **({"note": "consecutive integers; looks like an index"} if is_index else {}),
                }
            )
    return proposals


def command_prepare(args) -> int:
    workspace = Path(args.output).expanduser().resolve()
    files: list[Path] = []
    for item in args.input:
        path = Path(item).expanduser().resolve()
        if path.is_dir():
            files.extend(sorted(p for p in path.rglob("*") if p.suffix.lower() in TABLE_SUFFIXES))
        elif path.is_file():
            files.append(path)
        else:
            raise SystemExit(f"input not found: {path}")
    if not files:
        raise SystemExit("no CSV, TSV, or Excel files found")
    require_new_workspace(workspace)
    (workspace / "tables").mkdir()
    sources, series = [], []
    for path in files:
        tables = read_tables(path)
        for table, grid, formula_cells in tables:
            dump = workspace / "tables" / f"{slug(path.stem)}--{slug(table)}.csv"
            with dump.open("w", newline="", encoding="utf-8") as handle:
                csv.writer(handle).writerows(grid)
            series.extend(propose_series(path.name, table, grid, formula_cells))
        sources.append({"path": str(path), "tables": [name for name, _, _ in tables]})
    write_json(workspace / "series.json", {"sources": sources, "series": series, "means": [], "tests": []})
    print(
        json.dumps(
            {
                "workspace": str(workspace),
                "sources": len(sources),
                "proposed_series": len(series),
                "excluded_as_index": sum(not item["include"] for item in series),
            },
            indent=2,
        )
    )
    return 0


# --- scan ----------------------------------------------------------------


def command_scan(args) -> int:
    workspace = Path(args.workspace).expanduser().resolve()
    manifest = read_json(workspace / "series.json")
    series = [item for item in manifest.get("series", []) if item.get("include", True)]
    ids = [item["id"] for item in series]
    if len(ids) != len(set(ids)):
        raise ValueError("duplicate series ids in series.json")
    for item in series:
        cleaned = [clean_number(value) for value in item["values"]]
        if None in cleaned:
            raise ValueError(f"{item['id']} has a non-numeric value: {item['values'][cleaned.index(None)]!r}")
        item["values"] = cleaned

    findings: list[dict[str, Any]] = []

    def add(check: str, a: str, b: str = "", expected: bool = False, **result) -> None:
        findings.append({"check": check, "series_a": a, "series_b": b, "expected": expected, **result})

    pools: dict[str, list[dict]] = defaultdict(list)
    for item in series:
        pools[f"pooled:{item['source']}"].append(item)
        if result := terminal_digit(last_digits(item["values"])):
            add("terminal_digit", item["id"], **result)
        if result := decimal_repetition(item["values"]):
            add("decimal_repetition", item["id"], **result)
        if result := benford(item["values"]):
            add("benford", item["id"], **result)
    if len(pools) > 1:
        pools["pooled:all"] = series
    for name, members in pools.items():
        if len(members) < 2:
            continue  # identical to the per-series test
        if result := terminal_digit([digit for item in members for digit in last_digits(item["values"])]):
            add("terminal_digit", name, **result)
        if result := benford([value for item in members for value in item["values"]]):
            add("benford", name, **result)

    pairs_compared = 0
    for first, second in itertools.combinations(series, 2):
        a, b = first["values"], second["values"]
        if len(a) != len(b) or len(a) < args.min_rows or len(set(a)) < 3 or len(set(b)) < 3:
            continue
        pairs_compared += 1
        expected = bool(first.get("derivation_group")) and first.get("derivation_group") == second.get("derivation_group")
        formula = " [formula cells present]" if first.get("formula") or second.get("formula") else ""
        if relation := fixed_relation(a, b):
            add(
                f"fixed_{relation['relation']}",
                first["id"],
                second["id"],
                expected,
                n=len(a),
                flagged=True,
                detail=relation["detail"] + formula,
            )
        elif result := decimal_match(a, b):
            add("decimal_match", first["id"], second["id"], expected, **result)

    groups = {item["id"]: item.get("derivation_group") for item in series}
    for run in shared_runs(series, args.min_run):
        first, second = run.pop("series_a"), run.pop("series_b")
        expected = bool(groups[first]) and groups[first] == groups[second]
        add("shared_run", first, second, expected, flagged=True, **run)

    untestable = Counter()
    for entry in manifest.get("means", []):
        n, items = int(entry["n"]), int(entry.get("items", 1))
        result = grim(str(entry["mean"]), n, items, bool(entry.get("percent")))
        if result["testable"]:
            add("grim", entry["id"], n=n, flagged=not result["consistent"], detail=result["detail"])
        else:
            untestable["grim"] += 1
        if entry.get("sd") is None or entry.get("percent") or result.get("consistent") is False:
            continue  # a GRIM failure already explains the entry
        result = grimmer(str(entry["mean"]), str(entry["sd"]), n, items)
        if result["testable"]:
            add("grimmer", entry["id"], n=n, flagged=not result["consistent"], detail=result["detail"])
        else:
            untestable["grimmer"] += 1

    for entry in manifest.get("tests", []):
        result = statcheck(str(entry["report"]), int(entry.get("tails", 2)), args.alpha)
        add("statcheck", entry["id"], flagged=not result["consistent"], detail=result["detail"])

    tested = [row for row in findings if "p" in row]
    for row, q in zip(tested, benjamini_hochberg([row["p"] for row in tested])):
        row["q"] = q
        row["flagged"] = q < args.alpha and (row["check"] != "benford" or row["mad"] >= 0.015)
    # Keep only informative pairwise decimal matches; the rest still counted in the FDR family.
    findings = [
        row for row in findings if row["check"] != "decimal_match" or row["flagged"] or row["statistic"] >= 2
    ]
    findings.sort(key=lambda row: (not row["flagged"], row["expected"], row.get("q", 0.0)))

    write_json(workspace / "findings.json", findings)
    columns = ["check", "series_a", "series_b", "flagged", "expected", "n", "statistic", "p", "q", "detail"]
    with (workspace / "findings.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(findings)
    summary = {
        "series_in_scope": len(series),
        "series_excluded": len(manifest.get("series", [])) - len(series),
        "values_in_scope": sum(len(item["values"]) for item in series),
        "pairs_compared": pairs_compared,
        "p_value_tests": len(tested),
        "means_listed": len(manifest.get("means", [])),
        "untestable": dict(untestable),
        "tests_recomputed": len(manifest.get("tests", [])),
        "flagged_by_check": dict(Counter(row["check"] for row in findings if row["flagged"] and not row["expected"])),
        "expected_relationships": sum(row["flagged"] and row["expected"] for row in findings),
        "alpha_fdr": args.alpha,
    }
    write_json(workspace / "scan-summary.json", summary)
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)

    prepare = subparsers.add_parser("prepare", help="dump tables and propose numeric series")
    prepare.add_argument("--input", required=True, action="append", help="CSV/TSV/XLSX file or directory; may be repeated")
    prepare.add_argument("--output", required=True, help="new data-audit workspace")
    prepare.set_defaults(func=command_prepare)

    scan = subparsers.add_parser("scan", help="screen series.json")
    scan.add_argument("--workspace", required=True)
    scan.add_argument("--alpha", type=float, default=0.05, help="Benjamini-Hochberg FDR level")
    scan.add_argument("--min-rows", type=int, default=4, help="minimum paired rows for relation checks")
    scan.add_argument("--min-run", type=int, default=4, help="minimum repeated run length")
    scan.set_defaults(func=command_scan)
    return parser


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        return args.func(args)
    except (ValueError, KeyError, json.JSONDecodeError) as error:
        parser.error(str(error))
        return 2


if __name__ == "__main__":
    sys.exit(main())
