"""Editable, data-derived maintenance diagrams and native PowerPoint charts."""
from collections import Counter
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE,XL_LABEL_POSITION
from pptx.util import Inches,Pt
from mock_data import ROWS,EXPECTED,metrics

def add_visuals(n,g):
 def call(name,*a,**k):return g[name](*a,**k)
 def text(s,x,y,w,h,value,size=20,bold=False):return call('txt',s,Inches(x),Inches(y),Inches(w),Inches(h),[[(value,size,g['INK'],bold)]])
 def frame(title,kicker):s=call('slide');call('head',s,title,kicker=kicker);return s
 def finish(s):text(s,.85,6.55,11.7,.3,'Source: synthetic tblWorkOrders | 30 Sep 2026 | All filters All unless stated',11);call('footer',s)
 if n==1:
  call('flow_h','The Northstar Maintenance Evidence Chain',['Excel: clean and check','Word: operations brief','PowerPoint: decision story','Outlook: reviewed draft'],kicker='CONNECTED USE CASE',color=g['BLUE'])
  s=frame('A Work Order Has State, Scope and Ownership','EDITABLE DATA MODEL')
  cards=[('ID','WO-2026-0003'),('Site','Woodlands Hub'),('Category','Leak Repair'),('Status','In Progress'),('Owner','Chloe Goh')]
  for i,(k,v) in enumerate(cards):
   x=.85+(i%3)*4;y=2+(i//3)*1.6;call('rect',s,Inches(x),Inches(y),Inches(3.7),Inches(1.25),g['LIGHT']);text(s,x+.15,y+.1,3.4,.3,k,18,True);text(s,x+.15,y+.5,3.4,.6,v,23)
  text(s,.95,5.7,11.3,.5,'Current owner is evidence. Proposed owner is a separate planning output.',21);finish(s)
 elif n==6:
  s=frame('Anatomy of the Maintenance Skill','EDITABLE SKILL CONTRACT')
  fields=[('Discovery','name + description','When should Copilot use it?'),('Input guard','table or snapshot + date + scope','What evidence is missing?'),('Method','definitions + source preservation','How should measures be calculated?'),('Outputs','summary / charts / proposals','Which named areas may change?'),('Verification','counts + denominators + reruns','What proves the result?')]
  for i,(a,b,c) in enumerate(fields):
   y=1.8+i*.86;call('rect',s,Inches(.85),Inches(y),Inches(11.65),Inches(.72),g['LIGHT']);text(s,1,y+.12,2.2,.5,a,19,True);text(s,3.4,y+.12,4.5,.5,b,17);text(s,8,y+.12,4.2,.5,c,17)
  finish(s)
 if n==6:
  call('flow_h','A Skill Improves Through Evidence',['v1 instructions','Run on verified inputs','Find one concrete defect','v2 + repeat the check'],kicker='REPEATABLE IMPROVEMENT',color=g['TEAL'])
 elif n==7:
  s=frame('Northstar Portfolio at a Glance','WORKED DASHBOARD SNAPSHOT')
  data=[('Total orders',EXPECTED['total']),('Backlog',EXPECTED['backlog']),('Overdue',EXPECTED['overdue']),('Closed',EXPECTED['closed']),('On-time closed',f"{EXPECTED['on_time_closed']}/{EXPECTED['closed']}"),('Recorded cost',f"S${EXPECTED['cost']:,.0f}")]
  for i,(label,value) in enumerate(data):
   x=.85+i%3*4;y=1.9+i//3*1.95;call('rect',s,Inches(x),Inches(y),Inches(3.7),Inches(1.65),g['LIGHT']);text(s,x+.18,y+.12,3.4,.4,label,19,True);text(s,x+.18,y+.6,3.4,.7,str(value),35,True)
  finish(s)
  s=frame('Backlog by Site: Count the Nonterminal Work','NATIVE EDITABLE CHART')
  counts=Counter(r[1] for r in ROWS if r[4] not in ('Closed','Cancelled'));d=CategoryChartData();d.categories=list(counts);d.add_series('Backlog work orders',list(counts.values()));ch=s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED,Inches(.95),Inches(1.85),Inches(11.3),Inches(4.45),d).chart;ch.has_legend=False;ch.plots[0].has_data_labels=True;ch.plots[0].data_labels.position=XL_LABEL_POSITION.OUTSIDE_END;ch.plots[0].data_labels.font.size=Pt(16);ch.series[0].format.fill.solid();ch.series[0].format.fill.fore_color.rgb=g['TEAL'];ch.category_axis.tick_labels.font.size=Pt(16);ch.value_axis.tick_labels.font.size=Pt(14);finish(s)
 elif n==8:
  s=frame('Same Data, Different Filter Scope','EDITABLE WORKED COMPARISON')
  scopes=[('All sites',metrics()),('Woodlands Hub',metrics([r for r in ROWS if r[1]=='Woodlands Hub']))]
  for i,(label,e) in enumerate(scopes):
   x=.85+i*6;call('rect',s,Inches(x),Inches(1.85),Inches(5.65),Inches(4.35),g['LIGHT']);text(s,x+.2,2,5.2,.6,label,26,True)
   for j,(k,v) in enumerate([('Total',e['total']),('Backlog',e['backlog']),('Overdue',e['overdue']),('Closed',e['closed']),('On-time closed',f"{e['on_time_closed']}/{e['closed']}")]):text(s,x+.2,2.8+j*.56,5.2,.45,f'{k}: {v}',23)
  finish(s)
 elif n==9:
  call('flow_h','A Verified Chart Becomes a Decision',['Observed backlog','Operational implication','Proposed recovery action','Manager decision requested'],kicker='EVIDENCE TO NARRATIVE',color=g['VIOLET'])
 elif n==10:
  call('flow_h','The Maintenance Reporting Handoff',['Interactive workbook','Dated Word evidence','Static PPTX + PDF','Reviewed unsent email'],kicker='ARTIFACT LIFECYCLE',color=g['TEAL'])

 if n==3:
  call('flow_h','Four Scopes of Customisation',['Personalisation: your account','.Rules: this workbook','Skill: repeated method','Prompt: this request'],kicker='EDITABLE SCOPE COMPARISON',color=g['BLUE'])
 if n==4:
  s=frame('Visible .Rules: Instructions Travel with the Workbook','EDITABLE WORKBOOK MODEL')
  for i,(label,value) in enumerate([('A1','Northstar reporting conventions'),('A2','Preserve original IDs and current owners'),('A3','Backlog excludes Closed and Cancelled'),('A4','Reporting date: 2026-09-30')]):
   y=2+i*.85;call('rect',s,Inches(.9),Inches(y),Inches(11.5),Inches(.7),g['LIGHT']);text(s,1.1,y+.1,1.2,.45,label,20,True);text(s,2.5,y+.1,9.5,.45,value,21)
  text(s,1,5.7,11,.45,'Formula rule reads Control!B2; Copilot uses the recalculated result on the next request.',18);finish(s)
 if n==5:
  s=frame('Overdue Contribution: 50 Work Orders, Five Categories','NATIVE EDITABLE PARETO EVIDENCE')
  labels=['Pump Service','Meter Service','Leak Repair','Valve Repair','Pipe Replacement']
  for x,title,kind,values in [(.9,'Descending overdue counts',XL_CHART_TYPE.COLUMN_CLUSTERED,[18,12,8,8,4]),(6.8,'Cumulative share (%)',XL_CHART_TYPE.LINE,[36,60,76,92,100])]:
   d=CategoryChartData();d.categories=labels;d.add_series(title,values);ch=s.shapes.add_chart(kind,Inches(x),Inches(2),Inches(5.5),Inches(3.9),d).chart;ch.has_legend=False;ch.plots[0].has_data_labels=True;ch.plots[0].data_labels.font.size=Pt(12);ch.category_axis.tick_labels.font.size=Pt(11);ch.value_axis.tick_labels.font.size=Pt(12)
   if kind==XL_CHART_TYPE.LINE:ch.value_axis.minimum_scale=0;ch.value_axis.maximum_scale=100
  text(s,1,6,11,.4,'Four categories first reach 80%. Equal counts use alphabetical ties. Contribution is not cause.',17);finish(s)
 if n==7:
  call('flow_h','The Interactive Dashboard Skill Contract',['Named source + date','Six shared-scope KPIs','Four filters + three charts','Refresh + safe rerun checks'],kicker='REUSABLE DASHBOARD METHOD',color=g['TEAL'])
