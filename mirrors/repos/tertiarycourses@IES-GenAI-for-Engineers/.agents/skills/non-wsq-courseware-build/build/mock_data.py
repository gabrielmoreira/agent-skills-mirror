"""Deterministic synthetic maintenance inputs and templates."""
from pathlib import Path
from datetime import datetime,timedelta
from collections import Counter
import csv,json
from openpyxl import Workbook
from openpyxl.worksheet.table import Table,TableStyleInfo
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart,LineChart,Reference
from docx import Document
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from customisation_assets import PERSONAL_PREFS, RULES, WORK_ORDER_QA_SKILL, GENERIC_DASHBOARD_SKILL
import course_data as C
REPORT=datetime(2026,9,30)
SITES=['Woodlands Hub','Jurong Depot','Tampines Centre','Paya Lebar Office','Changi Workshop']
CATEGORIES=['Leak Repair','Valve Repair','Pipe Replacement','Pump Service','Meter Service']
STATUSES=['New','Assigned','In Progress','On Hold','Closed','Cancelled']
TECHS=['Aisha Tan','Ben Lim','Chloe Goh','Daniel Lee','Elena Ong','Farid Rahman','Grace Ng','Hari Kumar']
HEADERS=['Work_Order_ID','Site','Job_Category','Priority','Status','Assigned_To','Opened_Date','Due_Date','Closed_Date','Estimated_Hours','Actual_Hours','Cost_SGD','Description','Failure_Mode']
# Failure mode found on each job (engineering root-cause view). Counts per category sum to 24 and give a classic Pareto:
# Corrosion 42, Joint seal leak 30, Valve seat wear 22 = 78% of all 120 work orders.
FAILURE_MIX={'Leak Repair':[('Corrosion',12),('Joint seal leak',10),('Third-party damage',2)],
 'Pipe Replacement':[('Corrosion',22),('Pressure surge',2)],'Valve Repair':[('Valve seat wear',22),('Pressure surge',2)],
 'Pump Service':[('Pump bearing wear',10),('Joint seal leak',10),('Pressure surge',4)],
 'Meter Service':[('Corrosion',8),('Joint seal leak',10),('Meter fault',4),('Blockage',2)]}
_FM={c:[m for m,k in mix for _ in range(k)] for c,mix in FAILURE_MIX.items()};_seen=Counter()
ROWS=[]
for i in range(120):
 status='Closed' if i%5 in (0,1) else ('Cancelled' if i%10==9 else ['New','Assigned','In Progress','On Hold'][i%4])
 opened=REPORT-timedelta(days=5+i%80);due=opened+timedelta(days=7+i%10)
 closed=due+timedelta(days=(i//5)%5-2) if status=='Closed' else None
 # Keep every closure at or before the reporting date.
 if closed and closed>REPORT:closed=REPORT
 owner='' if status=='New' and i%3==0 else TECHS[i%8]
 cat=CATEGORIES[(i//4)%5];mode=_FM[cat][(_seen[cat]*7)%24];_seen[cat]+=1  # 7 is coprime with 24: spreads modes across the year
 ROWS.append([f'WO-2026-{i+1:04d}',SITES[(i//3)%5],CATEGORIES[(i//4)%5],['Low','Medium','High','Critical'][i%4],status,owner,opened,due,closed,2+i%7,(1+i%9) if status=='Closed' else 0,round(90+i%13*35.5,2),f'{CATEGORIES[(i//4)%5]} scheduled or reactive task {i+1:03d}',mode])
def metrics(rows=ROWS):
 closed=[r for r in rows if r[4]=='Closed'];back=[r for r in rows if r[4] not in ('Closed','Cancelled')]
 return {'total':len(rows),'closed':len(closed),'cancelled':sum(r[4]=='Cancelled' for r in rows),'backlog':len(back),'overdue':sum(r[7]<REPORT for r in back),'unassigned_backlog':sum(not r[5] for r in back),'on_time_closed':sum(r[8]<=r[7] for r in closed),'on_time_rate':sum(r[8]<=r[7] for r in closed)/len(closed),'cost':round(sum(r[11] for r in rows),2),'site_counts':dict(Counter(r[1] for r in rows)),'failure_counts':dict(Counter(r[13] for r in rows)),'category_counts':dict(Counter(r[2] for r in rows))}
EXPECTED=metrics()
def style(ws):
 ws.freeze_panes='A2'  # no sheet AutoFilter: on a table sheet Excel 'repairs' the file and drops the table
 for c in ws[1]:c.fill=PatternFill('solid',fgColor='1F6FEB');c.font=Font(color='FFFFFF',bold=True)
 for col in ws.columns:
  ws.column_dimensions[col[0].column_letter].width=min(45,max(17,len(str(col[0].value))+2))
 for row in ws:
  for c in row:
   c.alignment=Alignment(vertical='top',wrap_text=True)
   if isinstance(c.value,datetime):c.number_format='yyyy-mm-dd'
def table(ws,name):
 t=Table(displayName=name,ref=ws.dimensions);t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(t)
def add_dashboard(wb,ready=True):
 ws=wb.create_sheet('Dashboard');ws.append(['Northstar Maintenance Dashboard','Selection','As at','2026-09-30'])
 for label in ['Site','Job_Category','Status','Assigned_To']:ws.append([label,'All'])
 ws.append(['Measure','Current selection']); ws.append(['Total work orders']);ws.append(['Backlog']);ws.append(['Overdue backlog']);ws.append(['Closed']);ws.append(['On-time closure rate']);ws.append(['Cost SGD'])
 # SUMPRODUCT plus structured references includes rows appended to the table.
 mask='(IF($B$2="All",1,--(tblWorkOrders[Site]=$B$2)))*(IF($B$3="All",1,--(tblWorkOrders[Job_Category]=$B$3)))*(IF($B$4="All",1,--(tblWorkOrders[Status]=$B$4)))*(IF($B$5="All",1,--(tblWorkOrders[Assigned_To]=$B$5)))'
 formula={7:f'=SUMPRODUCT({mask}*--(tblWorkOrders[Work_Order_ID]<>""))',8:f'=SUMPRODUCT({mask}*--(tblWorkOrders[Status]<>"Closed")*--(tblWorkOrders[Status]<>"Cancelled"))',9:f'=SUMPRODUCT({mask}*--(tblWorkOrders[Status]<>"Closed")*--(tblWorkOrders[Status]<>"Cancelled")*--(tblWorkOrders[Due_Date]<Control!$B$2))',10:f'=SUMPRODUCT({mask}*--(tblWorkOrders[Status]="Closed"))',11:f'=IF(B10=0,"N/A",SUMPRODUCT({mask}*--(tblWorkOrders[Status]="Closed")*--(tblWorkOrders[Closed_Date]<=tblWorkOrders[Due_Date]))/B10)',12:f'=SUMPRODUCT({mask}*tblWorkOrders[Cost_SGD])'}
 lists=wb.create_sheet('Filter_Lists')
 for col,values in enumerate([SITES,CATEGORIES,STATUSES,TECHS+['']],1):
  for row,value in enumerate(['All']+values,1):lists.cell(row,col,value)
  from openpyxl.utils import get_column_letter
  letter=get_column_letter(col); dv=DataValidation(type='list',formula1=f"'Filter_Lists'!${letter}$1:${letter}${len(values)+1}");ws.add_data_validation(dv);dv.add(ws.cell(col+1,2))
 if ready:
  for row,f in formula.items():ws.cell(row,2,f)
  ws['B11'].number_format='0.0%';ws['B12'].number_format='"S$"#,##0.00'
  for start,title,col,vals in [(15,'Work orders by status','Status',STATUSES),(24,'Work orders by category','Job_Category',CATEGORIES),(33,'Work orders by site','Site',SITES)]:
   ws.cell(start,1,title);ws.cell(start,2,'Count')
   for k,value in enumerate(vals,start+1):ws.cell(k,1,value);ws.cell(k,2,f'=SUMPRODUCT({mask}*--(tblWorkOrders[{col}]=A{k}))')
   ch=BarChart();ch.title=title;ch.add_data(Reference(ws,min_col=2,min_row=start,max_row=start+len(vals)),titles_from_data=True);ch.set_categories(Reference(ws,min_col=1,min_row=start+1,max_row=start+len(vals)));ch.width=21;ch.height=8;ws.add_chart(ch,f'E{start-13}')
 style(ws);ws.column_dimensions['A'].width=30;ws.column_dimensions['B'].width=24
 lists.sheet_state='hidden'
def workbook(path,raw=False,dashboard=False,blank=False,rules=False):
 wb=Workbook();ws=wb.active;ws.title='About_This_Data';ws.append(['Field','Value']);ws.append(['Company','Northstar Water Services - fictional water-pipe maintenance and service company']);ws.append(['Reporting date',REPORT]);ws.append(['Scope','120 Q3 2026 work orders, five sites. All identities and records are synthetic.']);ws.append(['Meaning','Cost is cumulative recorded cost of all statuses; not a budget or productivity claim.']);ws.append(['Use','Synthetic training data. Work on this file in your lab folder.'])
 ws=wb.create_sheet('Work_Orders');ws.append(HEADERS)
 source=[list(r) for r in ROWS]
 if raw:
  source[3][2]=' Leak repair ';source[12][2]='Pump service ';source += [list(source[4]),list(source[19])]
 for r in source:ws.append(r)
 table(ws,'tblWorkOrders')
 ws=wb.create_sheet('Technicians');ws.append(['Technician','Primary_Category','Available_Hours_Week','Home_Site'])
 for i,t in enumerate(TECHS):ws.append([t,CATEGORIES[(i//4)%5],30 if i%2 else 35,SITES[i%5]])
 table(ws,'tblTechnicians')
 ws=wb.create_sheet('Control');ws.append(['Metric','Expected / control']);ws.append(['Report_Date',REPORT])
 for k,v in EXPECTED.items():
  if not isinstance(v,dict):ws.append([k,v])
 ws.append(['Backlog definition','Status is neither Closed nor Cancelled']);ws.append(['Overdue definition','Backlog and Due_Date strictly earlier than Report_Date']);ws.append(['On-time definition','Closed_Date <= Due_Date; denominator Closed only']);ws.append(['Unknown assignment','Blank Assigned_To stays blank; propose separately'])
 ws=wb.create_sheet('Prompt_Log');ws.append(['Step','App / mode','Prompt','What Copilot did','How I verified','Decision'])
 ws=wb.create_sheet('Checks');ws.append(['Check','Formula / evidence','Result']);ws.append(['Total clean rows','=ROWS(tblWorkOrders[Work_Order_ID])',''])
 # One learner log replaces the separate CSV registers and checklists.
 ws=wb.create_sheet('Evidence_Log');ws.append(['Lab','Step','Claim_or_Test','Expected','Observed','Source_and_Scope','As_At','Verified_By','Notes'])
 for ws in wb:style(ws)
 if rules:
  rules_ws=wb.create_sheet('.Rules')
  for rule in RULES:rules_ws.append([rule])
  rules_ws.append(['="Reporting date: "&TEXT(Control!B2,"yyyy-mm-dd")'])
  rules_ws.column_dimensions['A'].width=110
  for row in rules_ws:
   row[0].alignment=Alignment(wrap_text=True,vertical='top')
  rules_ws.sheet_state='visible'
 if dashboard:add_dashboard(wb,ready=not blank)
 path.parent.mkdir(parents=True,exist_ok=True);wb.save(path)
def doc(path,title,paras):
 d=Document();d.add_heading(title,0)
 for heading,body in paras:d.add_heading(heading,1);d.add_paragraph(body)
 d.save(path)
def deck(path,checkpoint=False,count=8):
 from pptx.chart.data import CategoryChartData
 from pptx.enum.chart import XL_CHART_TYPE,XL_LABEL_POSITION
 r=Presentation();r.slide_width=Inches(13.333);r.slide_height=Inches(7.5)
 titles=['Q3 maintenance: prioritise the overdue backlog','120 orders: reconcile the portfolio','Site backlog: compare open work','Status and category mix','30 of 48 closures were on time','6 backlog orders have no current owner','Proposed recovery actions and trade-offs','Decision requested: review the recovery proposal']
 if count==4:titles=['Northstar maintenance portfolio - 30 September','50 of 60 backlog orders are overdue','6 backlog orders need an assignment proposal','Decision requested: review the proposed priorities']
 bodies=[
 'Reporting date:30 September 2026. All sites/categories/statuses/owners. Synthetic training snapshot. Findings describe recorded work; actions below are proposed.',
 f"Closed {EXPECTED['closed']} + Cancelled {EXPECTED['cancelled']} + Backlog {EXPECTED['backlog']} = {EXPECTED['total']}. Recorded cost S${EXPECTED['cost']:,.2f}; this is not a budget or saving.",
 'Backlog excludes Closed and Cancelled. Compare counts across five sites before choosing recovery priorities. Frequency alone does not establish the cause.',
 'Status describes current workflow state; category describes job type. All-category counts sum to 120; category frequency is not a root-cause diagnosis.',
 f"On-time closure rate = {EXPECTED['on_time_closed']}/{EXPECTED['closed']} = {EXPECTED['on_time_rate']:.1%}. Overdue backlog: {EXPECTED['overdue']}, using Due_Date < 30 September 2026. Closed-only denominator, not all 120 orders.",
 f"Unassigned backlog: {EXPECTED['unassigned_backlog']}. Compare primary category and available hours across 8 synthetic technicians. Current owners remain unchanged; specific suitable/approved allocations UNKNOWN.",
 'Proposed: review the overdue queue by priority/site; validate constraints before suggesting owners; keep category/capacity evidence beside each proposal. Trade-off: backlog recovery may compete with newly opened work. Savings and approvals UNKNOWN.',
 'Requested decision: review and approve or revise the proposed recovery priorities and assignment checks. Decision owner, approval outcome and next review date UNKNOWN until supplied. No email sent and no source assignment changed.'
 ]
 if count==4:bodies=[bodies[1],bodies[4],bodies[5],bodies[7]]
 notes=[
 'Source tblWorkOrders, all filters All, as at 2026-09-30. State scope before interpretation. Ready-made example deck for this lab.',
 'Counts partition the selected source table exactly; recorded cost includes all statuses. Do not infer a budget or savings.',
 'Backlog chart counts Status neither Closed nor Cancelled. Site totals are source-derived and reconcile to 60. Counts alone do not explain causes.',
 'Status chart uses all 120 source rows; category comparison is nominal. Check units and avoid claims of causation.',
 'Closed on/before Due_Date is on time. Numerator 30 and denominator 48 are both source-derived. A different scope must be recalculated.',
 'Six blank current owners are in backlog. Technicians table contains planning inputs, not approval or comprehensive safety/capacity evidence.',
 'Actions are proposals linked to observed overdue/unassigned work. They need operational review; no promised saving or productivity gain.',
 'Ask for a decision without claiming one was made. The PowerPoint is a static snapshot; retain the interactive workbook and evidence register.'
 ]
 if count==4:notes=[notes[1],notes[4],notes[5],notes[7]]
 for i,title in enumerate(titles):
  sh=r.slides.add_slide(r.slide_layouts[6]);tb=sh.shapes.add_textbox(Inches(.7),Inches(.5),Inches(12),Inches(1));para=tb.text_frame.paragraphs[0];para.text=title;para.font.size=Pt(30);para.font.bold=True;para.font.color.rgb=RGBColor.from_string('0B3060')
  tb=sh.shapes.add_textbox(Inches(.8),Inches(1.65),Inches(11.7),Inches(1.5 if checkpoint else 3.6));tb.text_frame.word_wrap=True;para=tb.text_frame.paragraphs[0];para.text=bodies[i] if checkpoint else 'Add the verified chart or evidence table for this slide. Separate observed facts from proposed actions.';para.font.size=Pt(20)
  if checkpoint and ((count==8 and i in (2,3)) or (count==4 and i==1)):
   if count==8 and i==3:counts=Counter(row[4] for row in ROWS);label='All work orders by status'
   else:counts=Counter(row[1] for row in ROWS if row[4] not in ('Closed','Cancelled'));label='Backlog work orders by site'
   data=CategoryChartData();data.categories=list(counts);data.add_series(label,list(counts.values()));chart=sh.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED,Inches(1),Inches(3.3),Inches(11),Inches(3.2),data).chart;chart.has_legend=False;chart.plots[0].has_data_labels=True;chart.plots[0].data_labels.position=XL_LABEL_POSITION.OUTSIDE_END;chart.category_axis.tick_labels.font.size=Pt(12);chart.value_axis.tick_labels.font.size=Pt(12);chart.series[0].format.fill.solid();chart.series[0].format.fill.fore_color.rgb=RGBColor.from_string('13A89E')
  elif checkpoint:
   tb=sh.shapes.add_textbox(Inches(.8),Inches(3.6),Inches(11.7),Inches(2));para=tb.text_frame.paragraphs[0];para.text=['Observation -> evidence -> proposed action -> decision','Keep status, site, category and owner scope consistent.','Validate constraints before confirming an operational change.','Static presentation: verified snapshot, not an automatic live link.'][i%4];para.font.size=Pt(27);para.font.bold=True;para.font.color.rgb=RGBColor.from_string('1F6FEB')
  tb=sh.shapes.add_textbox(Inches(.8),Inches(6.8),Inches(11.7),Inches(.35));para=tb.text_frame.paragraphs[0];para.text=f'{C.TITLE} {C.VERSION} | Synthetic Northstar example | 2026-09-30 | All filters All | Slide{i+1}';para.font.size=Pt(10)
  sh.notes_slide.notes_text_frame.text=notes[i] if checkpoint else 'Provide the selected verified snapshot and replace scaffold content. Add source/date/scope notes; do not invent facts.'
 r.save(path)
def build_pareto_checkpoint(data,cp):
 from openpyxl import load_workbook
 counts=Counter(r[2] for r in ROWS if r[4] not in ('Closed','Cancelled') and r[7]<REPORT)
 ordered=sorted(counts.items(),key=lambda pair:(-pair[1],pair[0]));total=sum(counts.values());cum=0;controls=[]
 for label,count in ordered:
  cum+=count;controls.append(dict(Job_Category=label,Count=count,Cumulative_Count=cum,Share=count/total,Cumulative_Percent=cum/total))
 path=cp/'Northstar_Pareto_Checkpoint.xlsx';workbook(path,rules=True);wb=load_workbook(path);ws=wb.create_sheet('Pareto_Job_Category');ws.append(['Job_Category','Count','Cumulative_Count','Share','Cumulative_Percent'])
 for i,(label,count) in enumerate(ordered,2):
  ws.append([label,f'=COUNTIFS(tblWorkOrders[Job_Category],A{i},tblWorkOrders[Status],"<>Closed",tblWorkOrders[Status],"<>Cancelled",tblWorkOrders[Due_Date],"<"&Control!$B$2)',f'=SUM($B$2:B{i})',f'=IF(SUM($B$2:$B$6)=0,"N/A",B{i}/SUM($B$2:$B$6))',f'=IF(SUM($B$2:$B$6)=0,"N/A",C{i}/SUM($B$2:$B$6))'])
  ws.cell(i,4).number_format='0%';ws.cell(i,5).number_format='0%'
 ws['A8']='Reporting date';ws['B8']='=Control!B2';ws['B8'].number_format='dd-mmm-yyyy';ws['A9']='Scope';ws['B9']='All; overdue backlog only';ws['A10']='Interpretation';ws['B10']='Four categories first reach80%; counts show contribution, not cause.'
 bar=BarChart();bar.type='col';bar.title='Overdue backlog: category contribution';bar.y_axis.title='Work orders';bar.x_axis.title='Job category';bar.add_data(Reference(ws,min_col=2,min_row=1,max_row=6),titles_from_data=True);bar.set_categories(Reference(ws,min_col=1,min_row=2,max_row=6));bar.width=24;bar.height=13
 line=LineChart();line.add_data(Reference(ws,min_col=5,min_row=1,max_row=6),titles_from_data=True);line.set_categories(Reference(ws,min_col=1,min_row=2,max_row=6));line.y_axis.axId=200;line.y_axis.title='Cumulative share';line.y_axis.numFmt='0%';line.y_axis.scaling.min=0;line.y_axis.scaling.max=1;line.y_axis.crosses='max';
 bar+=line;ws.add_chart(bar,'G2');style(ws);ws.column_dimensions['A'].width=26;ws.column_dimensions['B'].width=24;wb.save(path)

REFRESH_ROWS=[['WO-2026-0121','Woodlands Hub','Leak Repair','High','Assigned','Aisha Tan',datetime(2026,9,25),datetime(2026,9,29),None,3,0,150,'Sandbox overdue test','Corrosion'],['WO-2026-0122','Woodlands Hub','Valve Repair','Medium','Closed','Ben Lim',datetime(2026,9,20),datetime(2026,9,30),datetime(2026,9,29),2,2,250,'Sandbox on-time test','Valve seat wear']]
def add_refresh_rows(path):
 from openpyxl import load_workbook
 wb=load_workbook(path);ws=wb.create_sheet('Refresh_Test_Rows');ws.append(HEADERS)
 for r in REFRESH_ROWS:ws.append(r)
 ws.append([]);ws.append(['Use only in the sandbox copy (Lab 8 Step 5): copy these two rows below the last row of tblWorkOrders. Do not add them to the baseline workbook.'])
 style(ws);wb.save(path)
def technicians():
 return [(t,CATEGORIES[(i//4)%5],30 if i%2 else 35,SITES[i%5]) for i,t in enumerate(TECHS)]
def handbook(path):
 """Synthetic knowledge file for the Lab 11 agent: definitions, sites, people and FAQ that match the data."""
 E=EXPECTED;d=Document();d.add_heading('Northstar Water Services - Service Handbook (synthetic)',0)
 d.add_paragraph('TRAINING DOCUMENT. Northstar Water Services is a fictional water-pipe maintenance and service company. Every site, person, number and policy in this handbook is synthetic and exists only for the GenAI for Engineers course. Reporting snapshot: 30 September 2026.')
 d.add_heading('1. Company and service hubs',1)
 d.add_paragraph('Northstar repairs and maintains water pipes, valves, pumps and meters for building owners. Work is run from five service hubs. Each hub has a hub lead who reviews proposals; only the Operations Manager approves reassignments.')
 t=d.add_table(rows=1,cols=3);t.style='Light Grid Accent 1';t.rows[0].cells[0].text,t.rows[0].cells[1].text,t.rows[0].cells[2].text='Hub','Hub lead (fictional)','Work orders in snapshot'
 for site,lead in zip(SITES,['M. Rahim','J. Koh','S. Pillai','L. Wong','R. Tan']):
  c=t.add_row().cells;c[0].text=site;c[1].text=lead;c[2].text=str(E['site_counts'][site])
 d.add_heading('2. Job categories',1)
 for cat,desc in zip(CATEGORIES,['Locate and fix leaks in supply or return pipes.','Repair or replace isolation and control valves.','Replace a damaged or corroded pipe section.','Service, repair or recommission pumps.','Install, test or replace water meters.']):d.add_paragraph(f'{cat}: {desc}',style='List Bullet')
 d.add_heading('3. Work-order statuses',1)
 for st,desc in zip(STATUSES,['Logged, not yet assigned or scheduled.','Owner assigned, work not started.','Work under way.','Paused (parts, access or customer constraint).','Work finished; Closed_Date recorded.','Withdrawn; no work required.']):d.add_paragraph(f'{st}: {desc}',style='List Bullet')
 d.add_heading('4. Definitions used in every report',1)
 for line in ['Backlog: Status is neither Closed nor Cancelled.','Overdue backlog: backlog with Due_Date strictly before the reporting date (30 September 2026). A job due on 30 September is not overdue.','On-time closure rate: Closed work orders with Closed_Date on or before Due_Date, divided by all Closed work orders. Show N/A when there are no Closed work orders.','Unassigned: Assigned_To is blank. Never invent an owner; propose one for a manager decision.','Recorded cost: the sum of Cost_SGD for the work orders in scope. It is not a budget, saving or productivity measure.']:d.add_paragraph(line,style='List Bullet')
 d.add_heading('5. Priorities',1)
 d.add_paragraph('Priority is Low, Medium, High or Critical. The Due_Date on each work order is the authoritative deadline; this handbook does not define separate response-time targets.')
 d.add_heading('6. Technicians (synthetic)',1)
 t=d.add_table(rows=1,cols=4);t.style='Light Grid Accent 1'
 for i,h in enumerate(['Technician','Primary category','Available hours per week','Home hub']):t.rows[0].cells[i].text=h
 for row in technicians():
  c=t.add_row().cells
  for i,v in enumerate(row):c[i].text=str(v)
 d.add_paragraph('Available hours are a planning input, not an approval or proof of suitability.')
 d.add_heading('7. Approvals and escalation',1)
 for line in ['Hub leads may propose an owner for unassigned backlog with category and available-hours evidence.','Only the Operations Manager approves an assignment or reassignment.','Assistants, agents and reports never approve, assign, close or edit work orders and never send messages.','Customer names, contact details and addresses are not held in these files; treat any such question as UNKNOWN.']:d.add_paragraph(line,style='List Bullet')
 d.add_heading('8. Snapshot facts at 30 September 2026',1)
 for line in [f"Total work orders: {E['total']} (Closed {E['closed']}, Cancelled {E['cancelled']}, Backlog {E['backlog']}).",f"Overdue backlog: {E['overdue']}. Unassigned backlog: {E['unassigned_backlog']}.",f"On-time closure: {E['on_time_closed']} of {E['closed']} Closed = {E['on_time_rate']:.1%}.",f"Recorded cost: S${E['cost']:,.2f}.",'Each hub and each job category has 24 work orders in this snapshot.']:d.add_paragraph(line,style='List Bullet')
 d.add_heading('9. Frequently asked questions',1)
 for q,a in [('Why is a job due today not overdue?','Overdue uses a strict before-date rule: Due_Date must be earlier than 30 September 2026.'),('Does cancelled work count in the backlog?','No. Backlog excludes Closed and Cancelled work orders.'),('Can the assistant reassign a technician?','No. It may describe a proposal and the evidence; the Operations Manager decides.'),('Where do the numbers come from?','From Northstar_Work_Orders_Clean.xlsx (table tblWorkOrders), the verified 120-row snapshot.')]:
  d.add_paragraph(q,style='List Number');d.add_paragraph(a)
 path.parent.mkdir(parents=True,exist_ok=True);d.save(path)
def clean_single_sheet(path):
 wb=Workbook();ws=wb.active;ws.title='Work_Orders';ws.append(HEADERS)
 for r in ROWS:ws.append(r)
 table(ws,'tblWorkOrders');style(ws);path.parent.mkdir(parents=True,exist_ok=True);wb.save(path)
def skill_zip(path,name,body):
 import zipfile
 path.parent.mkdir(parents=True,exist_ok=True)
 with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:z.writestr(f'{name}/SKILL.md',body)
def ops_brief_template(path):
 doc(path,'Northstar operations brief',[('Reporting scope','Record file, table, filters and reporting date.'),('Verified measures','Replace this paragraph with verified values and denominators.'),('Backlog and priority','Summarise evidence without guessing causal explanations.'),('Assignment proposals','Keep current assignment and proposed assignment separate.'),('Decision and open questions','Name the decision needed; use UNKNOWN for missing facts.')])
def snapshot(path):
 doc(path,'Northstar verified maintenance snapshot',[('Source and scope','Synthetic tblWorkOrders: 120 records; all filters All; 30 September 2026.'),('Verified controls',json.dumps(EXPECTED,indent=2)),('Interpretation','Values describe this fictional snapshot. Category mix is a count, not evidence of root cause. Proposals are unapproved.'),('Presentation outline','Portfolio; backlog; mix; SLA; capacity; actions; decision. Reconcile every figure against the dashboard before use.')])
LAB1_SOURCE='Northstar_Work_Orders.xlsx'
def lab1_workbook(path):
 """Completed Lab 1 workbook so Lab 2 can start without Lab 1: raw copy, Checks and Evidence_Log filled."""
 from openpyxl import load_workbook
 E=EXPECTED;wb=load_workbook(path)
 ws=wb['About_This_Data'];ws.append(['Lab 1 result','Completed Lab 1 workbook: 2 exact duplicate records removed (WO-2026-0005, WO-2026-0020) and 2 category labels standardised (WO-2026-0004, WO-2026-0013). Supplied so this lab can start straight away.'])
 raw=wb.create_sheet('Raw_Export',1);raw.append(HEADERS)
 source=[list(r) for r in ROWS];source[3][2]=' Leak repair ';source[12][2]='Pump service ';source+=[list(source[4]),list(source[19])]
 for r in source:raw.append(r)
 ctl={row[0].value:row[0].row for row in wb['Control'].iter_rows(min_row=2)}
 ws=wb['Checks'];ws.delete_rows(2,ws.max_row)
 back='tblWorkOrders[Status],"<>Closed",tblWorkOrders[Status],"<>Cancelled"'
 for label,key,f in [('Total clean rows','total','=ROWS(tblWorkOrders[Work_Order_ID])'),('Closed','closed','=COUNTIFS(tblWorkOrders[Status],"Closed")'),
   ('Cancelled','cancelled','=COUNTIFS(tblWorkOrders[Status],"Cancelled")'),('Backlog (neither Closed nor Cancelled)','backlog',f'=COUNTIFS({back})'),
   ('Overdue backlog (Due_Date before Report_Date)','overdue',f'=COUNTIFS({back},tblWorkOrders[Due_Date],"<"&Control!$B$2)'),
   ('Unassigned backlog (blank Assigned_To)','unassigned_backlog',f'=COUNTIFS({back},tblWorkOrders[Assigned_To],"")'),('Recorded cost SGD','cost','=SUM(tblWorkOrders[Cost_SGD])')]:
  r=ws.max_row+1;ws.append([label,f,f'=IF(B{r}=Control!B{ctl[key]},"PASS","FAIL")'])
 r=ws.max_row+1;ws.append(['Identity: total = closed + cancelled + backlog','=B3+B4+B5',f'=IF(B{r}=B2,"PASS","FAIL")'])
 ws.append(['Duplicate Work_Order_ID values','=SUMPRODUCT(--(COUNTIF(tblWorkOrders[Work_Order_ID],tblWorkOrders[Work_Order_ID])>1))',f'=IF(B{r+1}=0,"PASS","FAIL")'])
 ws['B8'].number_format='"S$"#,##0.00'
 ws=wb['Evidence_Log'];scope=f'{LAB1_SOURCE}, tblWorkOrders, all filters All'
 for step,claim,val in [(3,'Exact duplicate records removed','2 (WO-2026-0005, WO-2026-0020)'),(3,'Category labels standardised','2 (WO-2026-0004, WO-2026-0013)'),
   (4,'Total work orders',E['total']),(4,'Closed',E['closed']),(4,'Cancelled',E['cancelled']),(4,'Backlog',E['backlog']),(4,'Overdue backlog',E['overdue']),
   (4,'Unassigned backlog',E['unassigned_backlog']),(4,'Recorded cost SGD',f"S${E['cost']:,.2f}"),(4,'Identity total = closed + cancelled + backlog',f"{E['total']} = {E['closed']} + {E['cancelled']} + {E['backlog']}"),
   (6,'Word brief figures agree with Checks','All 8 figures match')]:
  ws.append(['Lab 1',f'Step {step}',claim,val,val,scope,REPORT,'Lab 1 model answer','PASS'])
 for row in wb['Raw_Export'].iter_rows(min_row=2):
  for c in row:
   if isinstance(c.value,datetime):c.number_format='yyyy-mm-dd'
 for name in ('About_This_Data','Raw_Export','Checks','Evidence_Log'):style(wb[name])
 wb['Checks'].column_dimensions['A'].width=44;wb['Checks'].column_dimensions['B'].width=60
 wb.save(path)
def lab1_brief(path):
 """Completed Lab 1 Word operations brief (same headings as the Lab 1 template), built from the verified controls."""
 E=EXPECTED;back=[r for r in ROWS if r[4] not in ('Closed','Cancelled')];over=[r for r in back if r[7]<REPORT]
 d=Document();d.add_heading('Northstar operations brief - 30 September 2026',0)
 d.add_paragraph('Finished operations brief supplied for this lab. Northstar Water Services is fictional; all records are synthetic.').runs[0].italic=True
 d.add_heading('Reporting scope',1)
 d.add_paragraph(f'Source: {LAB1_SOURCE}, table tblWorkOrders (Checks and Control sheets). Scope: all sites, job categories, statuses and owners. '
                 f'Reporting date: 30 September 2026. {E["total"]} unique work orders after removing 2 exact duplicate export records.')
 d.add_heading('Verified measures',1)
 t=d.add_table(rows=1,cols=3);t.style='Light Grid Accent 1'
 for i,h in enumerate(['Measure','Value','Definition']):t.rows[0].cells[i].text=h
 for row in [('Total work orders',E['total'],'Unique Work_Order_ID rows'),('Closed',E['closed'],'Status = Closed'),('Cancelled',E['cancelled'],'Status = Cancelled'),
   ('Backlog',E['backlog'],'Status is neither Closed nor Cancelled'),('Overdue backlog',E['overdue'],'Backlog with Due_Date before 30 September 2026'),
   ('Unassigned backlog',E['unassigned_backlog'],'Backlog with blank Assigned_To'),('Recorded cost',f"S${E['cost']:,.2f}",'Sum of Cost_SGD, all statuses; not a budget or saving')]:
  c=t.add_row().cells
  for i,v in enumerate(row):c[i].text=str(v)
 d.add_paragraph(f"Check: {E['total']} = {E['closed']} closed + {E['cancelled']} cancelled + {E['backlog']} backlog.")
 d.add_heading('Backlog and priority',1)
 site_b=Counter(r[1] for r in back);site_o=Counter(r[1] for r in over)
 d.add_paragraph(f"{E['overdue']} of {E['backlog']} backlog work orders are overdue. By site (backlog / overdue): "
                 + '; '.join(f'{s} {site_b[s]} / {site_o[s]}' for s in SITES) + '. These counts describe recorded work; causes are UNKNOWN.')
 d.add_heading('Assignment proposals',1)
 ids=', '.join(r[0] for r in back if not r[5])
 d.add_paragraph(f"Current: {E['unassigned_backlog']} backlog work orders have no owner ({ids}). Blank owners stay blank in the source.")
 d.add_paragraph('Proposed (unapproved): hub leads suggest an owner for each unassigned job using primary category and available hours; the Operations Manager decides.')
 d.add_heading('Decision and open questions',1)
 for line in [f"Proposed next action: review the {E['overdue']} overdue jobs by priority and site.",
              f"Proposed next action: prepare owner proposals for the {E['unassigned_backlog']} unassigned backlog jobs.",
              'Decision requested: the Operations Manager approves or revises the recovery priorities and owner proposals.',
              'Open questions: causes of overdue work, technician constraints, approval owner and review date are UNKNOWN.']:
  d.add_paragraph(line,style='List Bullet')
 d.save(path)
def plain_deck(path):
 """A deliberately plain 6-slide update (default template, black text, no colour or branding) for the deck-design skill lab."""
 E=EXPECTED;back=[r for r in ROWS if r[4] not in ('Closed','Cancelled')];site_o=Counter(r[1] for r in back if r[7]<REPORT)
 top=sorted(site_o.items(),key=lambda x:-x[1])[:3];fm=sorted(Counter(r[13] for r in ROWS).items(),key=lambda x:(-x[1],x[0]))[:3]
 r=Presentation()  # default 4:3 template: plain on purpose
 s=r.slides.add_slide(r.slide_layouts[0]);s.shapes.title.text='Northstar maintenance update';s.placeholders[1].text='Operations review - 30 September 2026'
 for title,lines in [('Work order portfolio',[f"{E['total']} work orders in Q3 2026",f"{E['closed']} closed, {E['cancelled']} cancelled, {E['backlog']} still open (backlog)",f"Recorded cost S${E['cost']:,.2f}"]),
   ('Backlog and overdue work',[f"{E['overdue']} of {E['backlog']} backlog jobs are overdue"]+[f'{k}: {v} overdue jobs' for k,v in top]),
   ('Top failure modes',[f'{k}: {v} work orders' for k,v in fm]+[f"These 3 failure modes cause {sum(v for _,v in fm)/E['total']:.0%} of all work orders"]),
   ('Unassigned jobs',[f"{E['unassigned_backlog']} backlog jobs have no owner",'Proposal: hub leads suggest owners by job category and free hours','No assignment is approved yet']),
   ('Decision needed',[f"Approve the recovery priorities for the {E['overdue']} overdue jobs",f"Approve or change the owner proposals for the {E['unassigned_backlog']} unassigned jobs",'Next review: UNKNOWN'])]:
  s=r.slides.add_slide(r.slide_layouts[1]);s.shapes.title.text=title;tf=s.placeholders[1].text_frame;tf.text=lines[0]
  for line in lines[1:]:tf.add_paragraph().text=line
 r.save(path)
def build_lab_data(num,folder):
 """Lean lab folder: the lab workbook and only the templates this lab uses (no checkpoints)."""
 root=Path(folder);data=root/'data';tmp=data/'templates';cp=data/'checkpoints'
 data.mkdir(parents=True,exist_ok=True)
 if num=='design':  # PowerPoint deck-design skill lab: a plain deck to transform plus the supplied skill
  plain_deck(data/'Northstar_Plain_Update.pptx');(tmp/'northstar-deck-design').mkdir(parents=True,exist_ok=True)
  from customisation_assets import DECK_DESIGN_SKILL
  (tmp/'northstar-deck-design/SKILL.md').write_text(DECK_DESIGN_SKILL,encoding='utf-8');return
 if num<=10:
  workbook(data/'Northstar_Work_Orders.xlsx',raw=(num==1),dashboard=(num>=8),rules=(num>=5))
  if num==8:add_refresh_rows(data/'Northstar_Work_Orders.xlsx')
  if num==2:  # Lab 1 results, named like Lab 1's outputs, for learners who missed Lab 1
   lab1_workbook(data/'Northstar_Work_Orders.xlsx')
   lab1_brief(data/'Northstar_Operations_Brief.docx')
 if num in (1,8):tmp.mkdir(parents=True,exist_ok=True);ops_brief_template(tmp/'Northstar_Operations_Brief_Template.docx')
 if num in (2,9,12):tmp.mkdir(parents=True,exist_ok=True);deck(tmp/'Northstar_Management_Deck_Template.pptx',count=4 if num==2 else 8)  # Lab 2 uses 4 slides
 if num==6:deck(data/'Northstar_Initial_Update.pptx',True,count=4)  # ready-made 4-slide update to improve
 if num==6:(tmp/'water-service-pareto').mkdir(parents=True,exist_ok=True);from customisation_assets import PARETO_SKILL;(tmp/'water-service-pareto/SKILL.md').write_text(PARETO_SKILL,encoding='utf-8')
 if num==10:deck(data/'Northstar_Management_Draft.pptx',True)  # ready-made 8-slide draft to polish
 if num==9:doc(tmp/'Northstar_Presentation_Brief_Template.docx','Northstar management presentation brief',[('Audience and decision','Audience: water service operations management. Decision: prioritise a proposed backlog recovery plan; approvals UNKNOWN.'),('Slide evidence','For each of eight slides specify headline, figure, chart, scope and source.'),('Limitations','Snapshot date, synthetic data, N/A for zero denominators, no invented savings.'),('Narrative','Observation -> operational implication -> proposed action -> decision request.')])
 if num==11:
  handbook(data/'knowledge/Northstar_Service_Handbook.docx');clean_single_sheet(data/'knowledge/Northstar_Work_Orders_Clean.xlsx')
  (tmp/'work-order-qa').mkdir(parents=True,exist_ok=True);(tmp/'work-order-qa/SKILL.md').write_text(WORK_ORDER_QA_SKILL,encoding='utf-8')
  # learners zip the folder themselves: ZIP files are never stored in the course repository
 if num==12:
  (tmp/'interactive-dashboard').mkdir(parents=True,exist_ok=True);(tmp/'interactive-dashboard/SKILL.md').write_text(GENERIC_DASHBOARD_SKILL,encoding='utf-8')
print('Northstar controls:',EXPECTED) if __name__=='__main__' else None
