"""Verify generated learner artifacts and independently recalculated dashboard outcomes."""
from pathlib import Path
import sys,json,re,html,zipfile
from openpyxl import load_workbook
from pptx import Presentation
import fitz
sys.path.insert(0,str(Path(__file__).parent))
from maintenance_labs import LABS
from build_prompts_page import split_prompts
from mock_data import ROWS,EXPECTED,metrics,HEADERS
import course_data as C
ROOT=Path(__file__).resolve().parents[4]
assert (ROOT/'courseware').is_dir(),ROOT
findings=[]
folders=sorted((ROOT/'labs').glob('lab-*'));assert len(folders)==10
for a,p in zip(LABS,folders):
 assert p.name.startswith(f'lab-{a["num"]:02d}-')
 for f in [f'Lab-{a["num"]:02d}-Instructions.pdf','data/Northstar_Work_Orders.xlsx','data/verified-controls.json','data/checkpoints/Northstar_Verified_Checkpoint.xlsx','data/checkpoints/Northstar_Verified_Snapshot.docx','data/templates/Northstar_Operations_Brief_Template.docx','data/templates/Northstar_Management_Deck_Template.pptx']:
  assert (p/f).exists(),p/f
 assert json.loads((p/'data/verified-controls.json').read_text())==EXPECTED
 wb=load_workbook(p/'data/Northstar_Work_Orders.xlsx');ws=wb['Work_Orders'];assert ws.max_row==123 if a['num']==1 else ws.max_row==121
 assert len({ws.cell(i,1).value for i in range(2,ws.max_row+1)})==120
 for row in ws.iter_rows(min_row=2):assert str(row[12].value).split(' scheduled')[0] == str(row[2].value).strip() or a['num']==1
 cmds=[cmd for _,cmd in a['steps'] if cmd]
 md=(ROOT/'docs/prompts/README.md').read_text();page=(ROOT/'docs/prompts/index.html').read_text()
 assert set(x.name for x in p.iterdir())=={'data',f'Lab-{a["num"]:02d}-Instructions.pdf'}
 for cmd in cmds:
  assert cmd in md
  for app,prompt in split_prompts(cmd):assert html.escape(prompt) in page and not re.match(r'^(EXCEL|WORD|POWERPOINT|OUTLOOK|MICROSOFT 365 COPILOT WEB)',prompt)
 assert [wb['Prompt_Log'].cell(i,3).value for i in range(2,2+len(cmds))]==cmds
 pdf=fitz.open(p/f'Lab-{a["num"]:02d}-Instructions.pdf');text=' '.join(pg.get_text() for pg in pdf);assert C.COURSE_CODE in text and C.VERSION in text
 assert any(link.get('uri')==f"{C.PROMPTS_PAGE}#lab-{a['num']}" for pg in pdf for link in pg.get_links()), 'Missing PDF prompt backlink'
 for instr,_ in a['steps']:
  for target in re.findall(r'data/[A-Za-z0-9_./-]+',instr):
   target=target.rstrip('.')
   assert (p/target).exists(),(p,target)
 if a['num']>=5:
  assert '.Rules' in wb.sheetnames and wb['.Rules'].sheet_state=='visible'
  assert wb['.Rules']['A1'].value
 findings.append({'lab':a['num'],'instruction_pdf_pages':len(pdf),'step_prompt_blocks':len(cmds),'copy_prompts':sum(len(split_prompts(cmd)) for cmd in cmds),'files':len([x for x in p.rglob('*') if x.is_file()])})
for label,selection in [('all',None),('site',lambda r:r[1]=='Woodlands Hub'),('empty',lambda r:r[1]=='Woodlands Hub' and r[2]=='Leak Repair' and r[5]=='Daniel Lee'),('zero_closed',lambda r:r[4]=='New')]:
 rows=ROWS if selection is None else [r for r in ROWS if selection(r)]
 e=metrics(rows) if any(r[4]=='Closed' for r in rows) else {'total':len(rows),'backlog':len([r for r in rows if r[4] not in ('Closed','Cancelled')]),'overdue':len([r for r in rows if r[4] not in ('Closed','Cancelled') and r[7]<__import__('mock_data').REPORT]),'closed':0,'on_time_rate':'N/A','cost':sum(r[11] for r in rows)}
 w=load_workbook(f'/tmp/c27-water-calc/out/{label}.xlsx',data_only=True)['Dashboard']
 actual=[w.cell(i,2).value for i in range(7,13)];expected=[e[k] for k in ['total','backlog','overdue','closed','on_time_rate','cost']];assert actual==expected,(label,actual,expected)
 for start,count in [(15,6),(24,5),(33,5)]:assert sum(w.cell(i,2).value for i in range(start+1,start+count+1))==e['total']
print('Verified four independently recalculated scopes, six KPIs and three chart totals each.')
w=load_workbook('/tmp/c27-water-calc/out/refresh.xlsx',data_only=True)['Dashboard'];actual=[w.cell(i,2).value for i in range(7,13)];expected=[122,61,51,49,31/49,36227.5]
for x,y in zip(actual,expected):assert abs(x-y)<1e-12,(actual,expected)
pareto=load_workbook('/tmp/c27-water-calc/out/pareto.xlsx',data_only=True)
assert [list(row) for row in list(pareto['Pareto_Job_Category'].values)[1:6]]==[['Pump Service',18,18,.36,.36],['Meter Service',12,30,.24,.6],['Leak Repair',8,38,.16,.76],['Valve Repair',8,46,.16,.92],['Pipe Replacement',4,50,.08,1]]
zero=load_workbook('/tmp/c27-water-calc/out/pareto_zero.xlsx',data_only=True)
assert all(zero['Pareto_Job_Category'].cell(row,4).value=='N/A' and zero['Pareto_Job_Category'].cell(row,5).value=='N/A' for row in range(2,7))
assert list(pareto['.Rules'].values)[-1][0]=='Reporting date: 2026-09-30'
checkpoint=next((ROOT/'labs').glob('lab-05-*/data/checkpoints/Northstar_Pareto_Checkpoint.xlsx'))
with zipfile.ZipFile(checkpoint) as z:
 chart=z.read('xl/charts/chart1.xml').decode();assert '<barChart>' in chart and '<lineChart>' in chart and 'axId val="200"' in chart and 'max val="1"' in chart
assert len([a for a in LABS if a['topic']==2])==4
assert all('@water-service-dashboard' in c for a in LABS if a['topic']==3 for _,c in a['steps'] if c.startswith('EXCEL - '))
print('Verified Pareto counts, cumulative shares, zero-total N/A, formula rule and native dual-axis chart.')
prs=Presentation(ROOT/'courseware'/f'{C.SHORT_TITLE}-{C.VERSION}.pptx');slide_count=len(prs.slides);pdf=fitz.open(ROOT/'courseware'/f'{C.SHORT_TITLE}-{C.VERSION}.pdf');assert len(pdf)==slide_count
concept=sum('CONCEPT' in '\n'.join(sh.text for sh in s.shapes if sh.has_text_frame) or 'HOW IT WORKS' in '\n'.join(sh.text for sh in s.shapes if sh.has_text_frame) or 'HOW TO USE IT' in '\n'.join(sh.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides)
assert concept>=48
schedule=C.SCHEDULE(lambda ns:', '.join(map(str,ns)))[1][1];assert sum(m for _,_,m,k,_ in schedule if k not in ('lunch','break'))==450
report={'course':C.COURSE_CODE,'version':C.VERSION,'slides':slide_count,'concept_slides_at_least':concept,'instructional_minutes':450,'scheduled_minutes_excluding_lunch':480,'baseline_controls':EXPECTED,'labs':findings,'dashboard_engine':'LibreOffice recalculation of actual XLSX formulas; all/site/empty/no-closed/append scopes verified','pareto_checks':'LibreOffice: 50 overdue, descending18/12/8/8/4, cumulative36/60/76/92/100, zero-total N/A; native combo chart XML verified','native_copilot_execution':'Not claimed; tenant access and skill availability must be confirmed by trainer.'}
(ROOT/f'docs/VALIDATION-{C.VERSION}.json').write_text(json.dumps(report,indent=2)+'\n');print(json.dumps(report,indent=2))
