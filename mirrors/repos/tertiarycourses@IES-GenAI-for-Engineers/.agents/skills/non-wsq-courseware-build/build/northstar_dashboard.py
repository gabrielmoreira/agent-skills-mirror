#!/usr/bin/env python3
"""Repair the Northstar lab workbooks and rebuild the interactive Dashboard.

1. Every workbook: drop worksheet AutoFilters and _xlnm._FilterDatabase names
   (a sheet AutoFilter on top of a table makes Excel "repair" the file and
   delete tblWorkOrders), and clean empty inlineStr / <v> cells.
2. Dashboard workbooks (except the Lab 7 starter): rebuild Dashboard with the
   SKILL contract (B2:B5 selectors, B7:B12 measures), KPI cards, three charts
   and four table slicers. Every formula respects both dropdowns and slicers.

Usage: python3 northstar_dashboard.py [labs_dir]  (build_labs.py calls process() on every lab workbook)
"""
import io, re, sys, zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.data_source import StrData, StrRef, StrVal
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

NAVY, BLUE, TEAL, GREY, PANEL, LINE, INPUT = "1F4E79", "2E75B6", "2A9D8F", "A6A6A6", "F3F6FA", "C9D3E3", "DDEBF7"
STATUSES = ["New", "Assigned", "In Progress", "On Hold", "Closed", "Cancelled"]
CATEGORIES = ["Leak Repair", "Valve Repair", "Pipe Replacement", "Pump Service", "Meter Service"]
SITES = ["Woodlands Hub", "Jurong Depot", "Tampines Centre", "Paya Lebar Office", "Changi Workshop"]
SLICERS = [("Site", 1), ("Job_Category", 1), ("Status", 2), ("Assigned_To", 2)]
from customisation_assets import RULES as RULES_LINES  # single source for the .Rules text
RULES_FORMULA = '="Reporting date: "&TEXT(Control!B2,"yyyy-mm-dd")'

SKIP_REBUILD = {"Northstar_Dashboard_Starter.xlsx"}

T = "tblWorkOrders"


def mask(p=""):
    """Rows selected by the Dashboard dropdowns (p = sheet prefix) AND visible after slicers."""
    return (f"(({T}[Site]={p}$B$2)+({p}$B$2=\"All\")>0)"
            f"*(({T}[Job_Category]={p}$B$3)+({p}$B$3=\"All\")>0)"
            f"*(({T}[Status]={p}$B$4)+({p}$B$4=\"All\")>0)"
            f"*(({T}[Assigned_To]={p}$B$5)+({p}$B$5=\"All\")+(({p}$B$5=\"Unassigned\")*({T}[Assigned_To]=\"\"))>0)"
            f"*SUBTOTAL(103,OFFSET({T}[[#Headers],[Work_Order_ID]],ROW({T}[Work_Order_ID])-ROW({T}[[#Headers],[Work_Order_ID]]),0))")


MASK, SMASK = mask(), mask("Dashboard!")
BACKLOG = f"*({T}[Status]<>\"Closed\")*({T}[Status]<>\"Cancelled\")"
OVERDUE = f"*({T}[Due_Date]<Control!$B$2)"
VISIBLE = f"SUBTOTAL(103,{T}[Work_Order_ID])"
ROWS = f"ROWS({T}[Work_Order_ID])"
SUM_SHEET = "Dashboard_Summaries"
S = f"{SUM_SHEET}!"


# ---------------------------------------------------------------- sanitise
def sanitise(data):
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            body = src.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet"):
                x = body.decode("utf-8")
                x = re.sub(r"<autoFilter[^>]*/>|<autoFilter[^>]*>.*?</autoFilter>", "", x)
                x = re.sub(r'<c ([^>]*?) t="inlineStr"\s*></c>', r"<c \1/>", x)
                x = x.replace("<v></v>", "")
                body = x.encode("utf-8")
            elif item.filename == "xl/workbook.xml":
                x = body.decode("utf-8")
                x = re.sub(r'<definedName name="_xlnm\._FilterDatabase"[^>]*>[^<]*</definedName>', "", x)
                x = x.replace("<definedNames></definedNames>", "")
                body = x.encode("utf-8")
            dst.writestr(item, body)
    return out.getvalue()


# ---------------------------------------------------------------- .Rules
def esc(t):
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")


def rewrite_rules(data):
    """Replace column A of a .Rules sheet with RULES_LINES plus the reporting-date formula."""
    src = zipfile.ZipFile(io.BytesIO(data))
    wbx = src.read("xl/workbook.xml").decode()
    m = re.search(r'<sheet [^>]*name="\.Rules"[^>]*r:id="([^"]+)"', wbx)
    if not m:
        return data
    rels = src.read("xl/_rels/workbook.xml.rels").decode()
    rel = next(r for r in re.findall(r"<Relationship [^>]+>", rels) if f'Id="{m.group(1)}"' in r)
    sheet = "xl/" + re.sub(r"^/?(xl/)?", "", re.search(r'Target="([^"]+)"', rel).group(1))
    rows = "".join(f'<row r="{i}"><c r="A{i}" s="2" t="inlineStr"><is><t>{esc(t)}</t></is></c></row>'
                   for i, t in enumerate(RULES_LINES, start=1))
    n = len(RULES_LINES) + 1
    rows += f'<row r="{n}"><c r="A{n}" s="2"><f>{esc(RULES_FORMULA[1:])}</f></c></row>'
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            body = src.read(item.filename)
            if item.filename == sheet:
                x = re.sub(r"<sheetData>.*?</sheetData>|<sheetData/>", f"<sheetData>{rows}</sheetData>", body.decode(), flags=re.S)
                body = re.sub(r'<dimension ref="[^"]*"/>', f'<dimension ref="A1:A{n}"/>', x).encode()
            dst.writestr(item, body)
    return out.getvalue()


# ---------------------------------------------------------------- dashboard
# Layout (BrightPath / C34 style): no merged cells, hidden gridlines, title in A1,
# KPI tiles in row 4, chart grid below, slicers on the right. A:B is the SKILL
# contract sidebar (B2:B5 dropdowns, B7:B12 measures).
TILE_COLS = "DEFGHI"
TILE_W, SIDE_W = 24, {"A": 19, "B": 17, "C": 2, "J": 2, "K": 31}
ROW_H = {1: 34, 2: 20, 3: 18, 4: 26, 5: 50, 6: 38, 7: 22, 8: 10}   # points
SLICER_COL = 10                                                    # column K (0-based)
SLICER_PX = [("Site", 160), ("Job_Category", 160), ("Status", 118), ("Assigned_To", 162)]


def col_px(width):
    return int(width * 7 + 5)


def row_top_px(row):
    """Pixel offset of the top of 1-based row."""
    return sum(round(ROW_H.get(r, 15) * 4 / 3) for r in range(1, row))


def px_to_row(y):
    """0-based row index and EMU offset for a pixel y."""
    r = 1
    while row_top_px(r + 1) <= y:
        r += 1
    return r - 1, (y - row_top_px(r)) * 9525


def side(c=LINE, s="thin"):
    return Side(style=s, color=c)


def put(ws, ref, value, **style):
    cell = ws[ref]
    cell.value = value
    for k, v in style.items():
        setattr(cell, k, v)
    return cell


def rich(size, bold=False, colour="404040"):
    cp = CharacterProperties(sz=size * 100, b=bold, solidFill=colour)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def chart(kind, title_ref, cats, series, horizontal, stacked=False, cached=""):
    ch = BarChart()
    ch.type = "bar" if horizontal else "col"
    cache = StrData(ptCount=1, pt=[StrVal(idx=0, v=cached)])
    ch.title = Title(tx=Text(strRef=StrRef(f=title_ref, strCache=cache)), overlay=False, txPr=rich(13, True))
    ch.gapWidth = 50
    for ref, colour in series:
        ch.add_data(ref, titles_from_data=True)
        ch.series[-1].graphicalProperties = GraphicalProperties(solidFill=colour)
    ch.set_categories(cats)
    ch.dataLabels = DataLabelList(numFmt="0;;;")
    ch.dataLabels.showVal = True
    for k in ("showSerName", "showCatName", "showLegendKey", "showPercent"):
        setattr(ch.dataLabels, k, False)
    ch.dataLabels.txPr = rich(9, colour="FFFFFF" if stacked else "404040")
    if stacked:
        ch.grouping, ch.overlap = "stacked", 100
        ch.legend.position = "b"
        ch.legend.txPr = rich(9)
    else:
        ch.legend = None
    ch.x_axis.delete = False
    ch.x_axis.txPr = rich(9)
    ch.y_axis.delete = True
    ch.y_axis.majorGridlines = None
    if horizontal:
        ch.x_axis.scaling.orientation = "maxMin"
    return ch


def build_summaries(wb, idx):
    """Formula-backed chart summaries, dynamic chart titles and KPI checks."""
    if SUM_SHEET in wb.sheetnames:
        del wb[SUM_SHEET]
    ss = wb.create_sheet(SUM_SHEET, idx)
    head = dict(font=Font(bold=True, color="FFFFFF"), fill=PatternFill("solid", fgColor=NAVY))
    put(ss, "A1", "Dashboard chart summaries, titles and checks (formula-backed; same selection as Dashboard)",
        font=Font(bold=True, size=13, color=NAVY))
    for col, w in {"A": 16, "B": 12, "C": 3, "D": 18, "E": 12, "F": 3, "G": 18, "H": 15, "I": 13, "J": 10,
                   "K": 11, "L": 9, "M": 3, "N": 30, "O": 70}.items():
        ss.column_dimensions[col].width = w

    for ref, text in {"A3": "Status", "B3": "Work orders", "D3": "Job category", "E3": "Work orders",
                      "G3": "Site", "H3": "Overdue backlog", "I3": "Other backlog", "J3": "Closed",
                      "K3": "Cancelled", "L3": "Total", "N3": "Item", "O3": "Value"}.items():
        put(ss, ref, text, **head)
    for i, s in enumerate(STATUSES, start=4):
        ss[f"A{i}"] = s
        ss[f"B{i}"] = f"=SUMPRODUCT({SMASK}*({T}[Status]=A{i}))"
    for i, c in enumerate(CATEGORIES, start=4):
        ss[f"D{i}"] = c
        ss[f"E{i}"] = f"=SUMPRODUCT({SMASK}*({T}[Job_Category]=D{i}))"
    for i, site in enumerate(SITES, start=4):
        on = f"*({T}[Site]=G{i})"
        ss[f"G{i}"] = site
        ss[f"H{i}"] = f"=SUMPRODUCT({SMASK}{BACKLOG}{OVERDUE}{on})"
        ss[f"I{i}"] = f"=SUMPRODUCT({SMASK}{BACKLOG}{on})-H{i}"
        ss[f"J{i}"] = f"=SUMPRODUCT({SMASK}*({T}[Status]=\"Closed\"){on})"
        ss[f"K{i}"] = f"=SUMPRODUCT({SMASK}*({T}[Status]=\"Cancelled\"){on})"
        ss[f"L{i}"] = f"=SUM(H{i}:K{i})"
    bold = Font(bold=True, color=NAVY)
    put(ss, "A10", "Total", font=bold); put(ss, "B10", "=SUM(B4:B9)", font=bold)
    put(ss, "D9", "Total", font=bold); put(ss, "E9", "=SUM(E4:E8)", font=bold)
    put(ss, "G9", "Total", font=bold)
    for col in "HIJKL":
        put(ss, f"{col}9", f"=SUM({col}4:{col}8)", font=bold)
    bad = dict(fill=PatternFill("solid", fgColor="FFC7CE"), font=Font(bold=True, color="9C0006"))
    for ref in ("B10", "E9", "L9"):
        ss.conditional_formatting.add(ref, FormulaRule(formula=[f"{ref}<>Dashboard!$B$7"], **bad))

    tie = lambda rng, what, noun: (f'IF(COUNTIF({rng},MAX({rng}))>1,COUNTIF({rng},MAX({rng}))&" {noun} tie '
                                   f'for most {what} ("&MAX({rng})&" each)"')
    items = [
        ("Status chart title",
         '=IF(Dashboard!$B$7=0,"No work orders in the current selection",'
         + tie("B4:B9", "work orders", "statuses")
         + ',INDEX(A4:A9,MATCH(MAX(B4:B9),B4:B9,0))&" is the largest status group ("&MAX(B4:B9)&" of "&Dashboard!$B$7&")"))'),
        ("Category chart title",
         '=IF(Dashboard!$B$7=0,"No work orders in the current selection",IF(MAX(E4:E8)=MIN(E4:E8),'
         '"All five job categories are level at "&MAX(E4:E8)&" work orders",'
         + tie("E4:E8", "work orders", "categories")
         + ',INDEX(D4:D8,MATCH(MAX(E4:E8),E4:E8,0))&" has the most work orders ("&MAX(E4:E8)&")")))'),
        ("Site chart title",
         '=IF(H9=0,"No overdue backlog in the current selection",'
         + tie("H4:H8", "overdue backlog", "sites")
         + ',INDEX(G4:G8,MATCH(MAX(H4:H8),H4:H8,0))&" has the most overdue backlog ("&MAX(H4:H8)&" of "&H9&")"))'),
        ("Unassigned backlog", f"=SUMPRODUCT({SMASK}{BACKLOG}*({T}[Assigned_To]=\"\"))"),
        ("Closed on or before due date",
         f"=SUMPRODUCT({SMASK}*({T}[Status]=\"Closed\")*({T}[Closed_Date]<={T}[Due_Date]))"),
        ("Baseline view (All, no slicer)",
         f'=AND(Dashboard!$B$2="All",Dashboard!$B$3="All",Dashboard!$B$4="All",Dashboard!$B$5="All",{VISIBLE}={ROWS})'),
        ("Charts reconcile to B7", "=AND(B10=Dashboard!$B$7,E9=Dashboard!$B$7,L9=Dashboard!$B$7)"),
        ("Check: total work orders",
         '=IF(AND(O10,OR(NOT(O9),Dashboard!$B$7=Control!$B$3)),"PASS","REVIEW")'),
        ("Check: backlog",
         '=IF(AND(Dashboard!$B$7=Dashboard!$B$8+Dashboard!$B$10+B9,OR(NOT(O9),Dashboard!$B$8=Control!$B$6)),"PASS","REVIEW")'),
        ("Check: overdue backlog",
         '=IF(AND(Dashboard!$B$9<=Dashboard!$B$8,H9=Dashboard!$B$9,OR(NOT(O9),Dashboard!$B$9=Control!$B$7)),"PASS","REVIEW")'),
        ("Check: closed", '=IF(AND(Dashboard!$B$10=B8,OR(NOT(O9),Dashboard!$B$10=Control!$B$4)),"PASS","REVIEW")'),
        ("Check: on-time rate",
         '=IF(Dashboard!$B$10=0,IF(Dashboard!$B$11="N/A","PASS","REVIEW"),IF(AND(ROUND(O8/Dashboard!$B$10-Dashboard!$B$11,6)=0,'
         'OR(NOT(O9),ROUND(Dashboard!$B$11-Control!$B$10,6)=0)),"PASS","REVIEW"))'),
        ("Check: recorded cost",
         f'=IF(AND(ROUND(Dashboard!$B$12-SUMPRODUCT({SMASK},{T}[Cost_SGD]),2)=0,'
         'OR(NOT(O9),ROUND(Dashboard!$B$12-Control!$B$11,2)=0)),"PASS","REVIEW")'),
    ]
    for r, (label, formula) in enumerate(items, start=4):
        ss[f"N{r}"] = label
        ss[f"O{r}"] = formula
    return ss


def baseline_titles(wb):
    """Text of the three dynamic chart titles at the All / no-slicer baseline (chart title cache)."""
    rows = list(wb["Work_Orders"].iter_rows(min_row=2, values_only=True))
    report = wb["Control"]["B2"].value
    status = {s: sum(r[4] == s for r in rows) for s in STATUSES}
    cats = {c: sum(r[2] == c for r in rows) for c in CATEGORIES}
    overdue = {t: sum(r[1] == t and r[4] not in ("Closed", "Cancelled") and r[7] < report for r in rows) for t in SITES}
    top = max(status, key=status.get)
    t1 = f"{top} is the largest status group ({status[top]} of {len(rows)})"
    if len(set(cats.values())) == 1:
        t2 = f"All five job categories are level at {max(cats.values())} work orders"
    else:
        best = max(cats, key=cats.get)
        t2 = f"{best} has the most work orders ({cats[best]})"
    worst = max(overdue, key=overdue.get)
    t3 = f"{worst} has the most overdue backlog ({overdue[worst]} of {sum(overdue.values())})"
    return t1, t2, t3


def build_dashboard(wb):
    lists = wb["Filter_Lists"]
    if "Unassigned" not in [c.value for c in lists["D"]]:
        lists.cell(row=max(c.row for c in lists["D"] if c.value) + 1, column=4, value="Unassigned")
    last = {col: max(c.row for c in lists[col] if c.value) for col in "ABCD"}

    idx = wb.sheetnames.index("Dashboard")
    del wb["Dashboard"]
    ws = wb.create_sheet("Dashboard", idx)
    ss = build_summaries(wb, idx + 1)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    for col, w in SIDE_W.items():
        ws.column_dimensions[col].width = w
    for col in TILE_COLS:
        ws.column_dimensions[col].width = TILE_W
    for r, h in ROW_H.items():
        ws.row_dimensions[r].height = h

    grey = Font(size=9, color="595959")
    panel = PatternFill("solid", fgColor=PANEL)
    mid = Alignment(vertical="center")

    # title (A1 overflows across the empty row)
    put(ws, "A1", '="Northstar Maintenance Dashboard — Q3 work orders to "&TEXT(Control!$B$2,"dd-mmm-yyyy")',
        font=Font(bold=True, size=20, color=NAVY), alignment=mid)
    put(ws, "D2", '="Data as at "&TEXT(Control!$B$2,"yyyy-mm-dd")&"   ·   Site: "&$B$2&"  ·  Category: "&$B$3'
                  '&"  ·  Status: "&$B$4&"  ·  Owner: "&$B$5',
        font=Font(italic=True, size=11, color="404040"), alignment=mid)
    put(ws, "D3", f'=IF({VISIBLE}<{ROWS},"⚠ Slicer filter active: "&{VISIBLE}&" of "&{ROWS}'
                  f'&" source rows included. Clear the slicers for the full baseline.",'
                  f'"All "&{ROWS}&" source rows available (no slicer filter).")',
        font=grey, alignment=mid)
    ws.conditional_formatting.add("D3", FormulaRule(formula=['LEFT($D$3,1)="⚠"'],
                                                    font=Font(bold=True, color="C00000")))

    # sidebar: SKILL contract
    for r in range(2, 14):
        for col in "AB":
            ws[f"{col}{r}"].fill = panel
            ws[f"{col}{r}"].alignment = mid
    for r, label, col in zip(range(2, 6), ["Site", "Job_Category", "Status", "Assigned_To"], "ABCD"):
        put(ws, f"A{r}", label, font=Font(bold=True, color=NAVY))
        put(ws, f"B{r}", "All", font=Font(bold=True, color=NAVY), fill=PatternFill("solid", fgColor=INPUT),
            border=Border(*(side(BLUE),) * 4))
        dv = DataValidation(type="list", formula1=f"'Filter_Lists'!${col}$1:${col}${last[col]}", allow_blank=False,
                            showErrorMessage=True, errorTitle="Choose from the list",
                            error="Pick All or a value from the dropdown.")
        dv.add(f"B{r}")
        ws.add_data_validation(dv)
    head = dict(font=Font(bold=True, color="FFFFFF"), fill=PatternFill("solid", fgColor=NAVY))
    put(ws, "A6", "Measure", **head)
    put(ws, "B6", "Current selection", **head)
    measures = [
        ("Total work orders", f"=SUMPRODUCT({MASK})", "0"),
        ("Backlog", f"=SUMPRODUCT({MASK}{BACKLOG})", "0"),
        ("Overdue backlog", f"=SUMPRODUCT({MASK}{BACKLOG}{OVERDUE})", "0"),
        ("Closed", f"=SUMPRODUCT({MASK}*({T}[Status]=\"Closed\"))", "0"),
        ("On-time closure rate", f"=IF(B10=0,\"N/A\",{S}O8/B10)", "0.0%"),
        ("Recorded cost (S$)", f"=SUMPRODUCT({MASK},{T}[Cost_SGD])", '"S$"#,##0.00'),
    ]
    for r, (label, formula, fmt) in zip(range(7, 13), measures):
        put(ws, f"A{r}", label, border=Border(bottom=side()))
        put(ws, f"B{r}", formula, number_format=fmt, font=Font(bold=True, color=NAVY),
            alignment=Alignment(horizontal="right", vertical="center"), border=Border(bottom=side()))
    put(ws, "A13", "Rows after slicers", font=Font(italic=True, color="595959"))
    put(ws, "B13", f"={VISIBLE}", font=Font(italic=True, color="595959"),
        alignment=Alignment(horizontal="right", vertical="center"))
    notes = ["How to filter", "• Dropdowns B2:B5 or the slicers", "  both filter every tile and chart.",
             "• Unassigned = blank Assigned_To.", "• Baseline: all All, slicers cleared.",
             "• Summaries + checks:", "  Dashboard_Summaries sheet."]
    for r, text in enumerate(notes, start=15):
        put(ws, f"A{r}", text, font=Font(bold=(r == 15), size=9, color=NAVY if r == 15 else "595959"))

    # KPI tiles (one column each, row 4 title, 5 value, 6 caption, 7 checks)
    tiles = [
        ("Total work orders", "=B7", "0", NAVY, f'="Selected of "&{ROWS}&" rows in tblWorkOrders"', "O11"),
        ("Backlog", "=B8", "0", BLUE,
         f'=IF(B7=0,"No rows selected",TEXT(B8/B7,"0%")&" of selection; "&{S}O7&" unassigned")', "O12"),
        ("Overdue backlog", "=B9", "0", BLUE, '="Backlog due before "&TEXT(Control!$B$2,"dd-mmm-yyyy")', "O13"),
        ("Closed", "=B10", "0", TEAL, f'="Cancelled in selection: "&{S}B9', "O14"),
        ("On-time closure rate", "=B11", "0.0%", NAVY,
         f'=IF(B10=0,"No closed jobs in selection",{S}O8&" of "&B10&" closed by due date")', "O15"),
        ("Recorded cost", "=B12", '"S$"#,##0.00', TEAL, '="Cost_SGD, all statuses; not a budget"', "O16"),
    ]
    gap = side("FFFFFF", "thick")
    for col, (title, value, fmt, colour, caption, check) in zip(TILE_COLS, tiles):
        fill = PatternFill("solid", fgColor=colour)
        centre = lambda wrap=False: Alignment(horizontal="center", vertical="center", wrap_text=wrap)
        put(ws, f"{col}4", title, font=Font(bold=True, size=12, color="FFFFFF"), alignment=centre())
        put(ws, f"{col}5", value, number_format=fmt, font=Font(bold=True, size=28, color="FFFFFF"),
            alignment=Alignment(horizontal="center", vertical="center", shrink_to_fit=True))
        put(ws, f"{col}6", caption, font=Font(size=9, color="FFFFFF"), alignment=centre(True))
        put(ws, f"{col}7", f'="Checks: "&{S}{check}', font=Font(bold=True, size=9, color="FFFFFF"),
            alignment=centre())
        for r in range(4, 8):
            ws[f"{col}{r}"].fill = fill
            ws[f"{col}{r}"].border = Border(left=gap, right=gap)

    # charts: 2 across, then full-width stacked site chart
    tile_cm = col_px(TILE_W) / 37.8
    t1, t2, t3 = baseline_titles(wb)
    top = Reference(ss, min_col=1, min_row=4, max_row=9)
    ch = chart("col", f"'{SUM_SHEET}'!$O$4", top,
               [(Reference(ss, min_col=2, min_row=3, max_row=9), NAVY)], horizontal=False, cached=t1)
    ch.width, ch.height = tile_cm * 3, 8.2
    ws.add_chart(ch, "D9")
    ch = chart("bar", f"'{SUM_SHEET}'!$O$5", Reference(ss, min_col=4, min_row=4, max_row=8),
               [(Reference(ss, min_col=5, min_row=3, max_row=8), BLUE)], horizontal=True, cached=t2)
    ch.width, ch.height = tile_cm * 3, 8.2
    ws.add_chart(ch, "G9")
    ch = chart("bar", f"'{SUM_SHEET}'!$O$6", Reference(ss, min_col=7, min_row=4, max_row=8),
               [(Reference(ss, min_col=c, min_row=3, max_row=8), colour)
                for c, colour in zip(range(8, 12), [NAVY, BLUE, TEAL, GREY])], horizontal=True, stacked=True, cached=t3)
    ch.width, ch.height = tile_cm * 6, 8.6
    ws.add_chart(ch, "D26")
    return ws


# ---------------------------------------------------------------- slicers
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def anchor(name, shape_id, y1, y2):
    (r1, o1), (r2, o2) = px_to_row(y1), px_to_row(y2)
    w = (col_px(SIDE_W["K"]) - 8) * 9525
    return (
        f'<xdr:twoCellAnchor editAs="oneCell"><xdr:from><xdr:col>{SLICER_COL}</xdr:col><xdr:colOff>38100</xdr:colOff>'
        f'<xdr:row>{r1}</xdr:row><xdr:rowOff>{o1}</xdr:rowOff></xdr:from><xdr:to><xdr:col>{SLICER_COL}</xdr:col>'
        f'<xdr:colOff>{w}</xdr:colOff><xdr:row>{r2}</xdr:row><xdr:rowOff>{o2}</xdr:rowOff></xdr:to>'
        '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
        '<mc:Choice xmlns:sle15="http://schemas.microsoft.com/office/drawing/2012/slicer" Requires="sle15">'
        f'<xdr:graphicFrame macro=""><xdr:nvGraphicFramePr><xdr:cNvPr id="{shape_id}" name="{name}"/>'
        '<xdr:cNvGraphicFramePr/></xdr:nvGraphicFramePr><xdr:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/></xdr:xfrm>'
        '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/drawing/2010/slicer">'
        f'<sle:slicer xmlns:sle="http://schemas.microsoft.com/office/drawing/2010/slicer" name="{name}"/>'
        '</a:graphicData></a:graphic></xdr:graphicFrame></mc:Choice><mc:Fallback>'
        '<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="0" name=""/><xdr:cNvSpPr><a:spLocks noTextEdit="1"/>'
        '</xdr:cNvSpPr></xdr:nvSpPr><xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="2300000" cy="1600000"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:prstClr val="white"/></a:solidFill>'
        '<a:ln w="1"><a:solidFill><a:prstClr val="green"/></a:solidFill></a:ln></xdr:spPr><xdr:txBody>'
        '<a:bodyPr vertOverflow="clip" horzOverflow="clip"/><a:lstStyle/><a:p><a:r><a:rPr lang="en-US" sz="1100"/>'
        '<a:t>This shape represents a table slicer. Table slicers need Excel 2013 or later.</a:t></a:r></a:p>'
        '</xdr:txBody></xdr:sp></mc:Fallback></mc:AlternateContent><xdr:clientData/></xdr:twoCellAnchor>')


def add_slicers(data):
    src = zipfile.ZipFile(io.BytesIO(data))
    files = {n: src.read(n) for n in src.namelist()}
    text = lambda n: files[n].decode("utf-8")

    wb_rels = text("xl/_rels/workbook.xml.rels")
    wbx = text("xl/workbook.xml")
    rid = re.search(r'<sheet [^>]*name="Dashboard"[^>]*r:id="([^"]+)"', wbx).group(1)
    rel = next(r for r in re.findall(r"<Relationship [^>]+>", wb_rels) if f'Id="{rid}"' in r)
    sheet = "xl/" + re.sub(r"^/?(xl/)?", "", re.search(r'Target="([^"]+)"', rel).group(1))
    sheet_rels = sheet.replace("worksheets/", "worksheets/_rels/") + ".rels"
    drawing = "xl/" + re.search(r'Target="(?:\.\./|/xl/)([^"]*drawings/[^"]+)"', text(sheet_rels)).group(1)

    table = next(n for n in files if n.startswith("xl/tables/") and f'name="{T}"' in text(n))
    tx = text(table)
    table_id = re.search(r'<table[^>]* id="(\d+)"', tx).group(1)
    col_ids = dict((name, cid) for cid, name in re.findall(r'<tableColumn id="(\d+)" name="([^"]+)"', tx))

    ct = text("[Content_Types].xml")
    names, caches, wb_ids = [], [], []
    for i, (col, ncols) in enumerate(SLICERS, start=1):
        files[f"xl/slicerCaches/slicerCache{i}.xml"] = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<slicerCacheDefinition xmlns="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main" '
            'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" mc:Ignorable="x" '
            'xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            f'name="Slicer_{col}" sourceName="{col}"><extLst><x:ext uri="{{2F2917AC-EB37-4324-AD4E-5DD8C200BD13}}" '
            'xmlns:x15="http://schemas.microsoft.com/office/spreadsheetml/2010/11/main">'
            f'<x15:tableSlicerCache tableId="{table_id}" column="{col_ids[col]}"/></x:ext></extLst>'
            '</slicerCacheDefinition>').encode()
        ct = ct.replace("</Types>", f'<Override PartName="/xl/slicerCaches/slicerCache{i}.xml" '
                                    'ContentType="application/vnd.ms-excel.slicerCache+xml"/></Types>')
        wb_rid = f"rIdSlicerCache{i}"
        wb_rels = wb_rels.replace("</Relationships>", f'<Relationship Id="{wb_rid}" '
                                  'Type="http://schemas.microsoft.com/office/2007/relationships/slicerCache" '
                                  f'Target="slicerCaches/slicerCache{i}.xml"/></Relationships>')
        wb_ids.append(wb_rid)
        names.append(f'<definedName name="Slicer_{col}">#N/A</definedName>')
        caches.append(f'<slicer name="{col}" cache="Slicer_{col}" caption="{col.replace("_", " ")}" '
                      f'columnCount="{ncols}" style="SlicerStyleLight1" rowHeight="225425"/>')

    files["xl/slicers/slicer1.xml"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<slicers xmlns="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" mc:Ignorable="x" '
        'xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main">' + "".join(caches) + "</slicers>").encode()
    ct = ct.replace("</Types>", '<Override PartName="/xl/slicers/slicer1.xml" '
                                'ContentType="application/vnd.ms-excel.slicer+xml"/></Types>')

    # workbook: defined names + slicer cache list
    wbx = wbx.replace("<definedNames/>", "")
    if "<definedNames>" in wbx:
        wbx = wbx.replace("</definedNames>", "".join(names) + "</definedNames>")
    else:
        wbx = wbx.replace("</sheets>", "</sheets><definedNames>" + "".join(names) + "</definedNames>")
    ext = ('<extLst><ext uri="{46BE6895-7355-4a93-B00E-2C351335B9C9}" '
           'xmlns:x15="http://schemas.microsoft.com/office/spreadsheetml/2010/11/main">'
           '<x15:slicerCaches xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
           + "".join(f'<x14:slicerCache xmlns:r="{NS_R}" r:id="{r}"/>' for r in wb_ids)
           + "</x15:slicerCaches></ext></extLst>")
    wbx = wbx.replace("</workbook>", ext + "</workbook>")

    # sheet: slicer relationship + slicerList
    files[sheet_rels] = text(sheet_rels).replace(
        "</Relationships>", '<Relationship Id="rIdSlicer1" Type="http://schemas.microsoft.com/office/2007/relationships/slicer" '
                            'Target="../slicers/slicer1.xml"/></Relationships>').encode()
    sx = text(sheet)
    sx = sx.replace("</worksheet>", '<extLst><ext uri="{3A4CF648-6AED-40f4-86FF-DC5316D8AED3}" '
                    'xmlns:x15="http://schemas.microsoft.com/office/spreadsheetml/2010/11/main">'
                    '<x14:slicerList xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
                    f'<x14:slicer xmlns:r="{NS_R}" r:id="rIdSlicer1"/></x14:slicerList></ext></extLst></worksheet>')

    # drawing: slicer frames (normalise to xdr/a prefixes)
    dx = text(drawing)
    if "xmlns:xdr=" not in dx:
        dx = dx.replace('<wsDr xmlns="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"',
                        '<wsDr xmlns="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
                        'xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"', 1)
    if 'xmlns:a="' not in dx.split(">", 1)[0]:
        dx = re.sub(r"(<(?:xdr:)?wsDr\b)", r'\1 xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"', dx, count=1)
    frames, y = "", row_top_px(4)
    for i, (col, h) in enumerate(SLICER_PX):
        frames += anchor(col, 100 + i, y, y + h)
        y += h + 10
    dx = re.sub(r"(</(?:xdr:)?wsDr>)", frames + r"\1", dx)

    for n in [n for n in files if n.startswith("xl/charts/chart")]:
        cx = text(n).replace('<numFmt formatCode="0;;;"/>', '<numFmt formatCode="0;;;" sourceLinked="0"/>')
        files[n] = re.sub(r"(<txPr>.*?)<a:r><a:t></a:t></a:r>", r"\1", cx, flags=re.S).encode()

    files.update({"xl/workbook.xml": wbx.encode(), "xl/_rels/workbook.xml.rels": wb_rels.encode(),
                  sheet: sx.encode(), drawing: dx.encode(), "[Content_Types].xml": ct.encode()})
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for n in ["[Content_Types].xml"] + [n for n in files if n != "[Content_Types].xml"]:
            dst.writestr(n, files[n])
    return out.getvalue()


# ---------------------------------------------------------------- main
def process(path):
    data = rewrite_rules(sanitise(path.read_bytes()))
    rebuilt = False
    if b'name="Dashboard"' in zipfile.ZipFile(io.BytesIO(data)).read("xl/workbook.xml") and path.name not in SKIP_REBUILD:
        wb = load_workbook(io.BytesIO(data))
        build_dashboard(wb)
        buf = io.BytesIO()
        wb.save(buf)
        data = add_slicers(buf.getvalue())
        rebuilt = True
    path.write_bytes(data)
    return rebuilt


if __name__ == "__main__":
    labs = Path(sys.argv[1] if len(sys.argv) > 1 else "labs")
    for p in sorted(labs.rglob("*.xlsx")):
        if p.name.startswith("~$"):
            continue
        print(("REBUILT  " if process(p) else "cleaned  ") + str(p))
