"""Reconcile the static TOC with actual final PDF pages, including multi-page TOCs."""
from pathlib import Path
import re,sys
from docx import Document
import fitz
norm=lambda t:re.sub('[^a-z0-9]','',t.lower())
path=Path(sys.argv[1]);pdf=Path(sys.argv[2]);doc=Document(path)
heads=[(p.text,int(p.style.name.split()[-1])) for p in doc.paragraphs if p.style.name.startswith('Heading') and p.style.name.split()[-1].isdigit() and int(p.style.name.split()[-1])<=2]
pages=[[norm(l) for l in page.get_text().splitlines()] for page in fitz.open(pdf)]
def starts(full,lines):
 for i,line in enumerate(lines):
  if len(line)>=min(8,len(full)) and full.startswith(line) and ''.join(lines[i:i+3]).startswith(full):return True
 return False
cursor=next(i for i,lines in enumerate(pages) if starts(norm(heads[0][0]),lines))
actual={}
for text,level in heads:
 target=norm(text)
 match=next((i for i in range(cursor,len(pages)) if starts(target,pages[i])),None)
 assert match is not None, text
 actual[text]=match+1;cursor=match
changed=0
for p in doc.paragraphs:
 if '\t' not in p.text:continue
 title,_,number=p.text.rpartition('\t')
 if title in actual and number.isdigit() and int(number)!=actual[title]:
  p.runs[0].text=title+'\t'+str(actual[title])
  for r in p.runs[1:]:r.text=''
  changed+=1
if changed:doc.save(path)
print(f'Actual-page TOC reconciliation: {path.name}, {len(actual)} entries, {changed} corrected.')
