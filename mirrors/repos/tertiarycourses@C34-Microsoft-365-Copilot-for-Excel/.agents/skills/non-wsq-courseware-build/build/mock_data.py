"""Deterministic synthetic HR / recruitment datasets for the C34 labs.

Scenario: BrightPath Talent Solutions — a FICTIONAL Singapore recruitment,
flexible-staffing and HR-outsourcing firm. Every company, person, ID and value
in these workbooks is synthetic and generated from fixed random seeds, so the
expected answers quoted in each lab's "Test it" step are reproducible.

`build_lab_data(num, folder)` writes the lab's workbook(s) into <folder>/data/.
`EXPECTED[num]` holds the values the lab text quotes for verification.
"""
import os
import random
import statistics
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

COMPANY = "BrightPath Talent Solutions"
FICTION_NOTE = (
    "All organisations, people, IDs and values in this workbook are FICTIONAL training data "
    "created for the Microsoft 365 Copilot for Excel course (C34). Any resemblance to a real "
    "company or person is coincidental."
)
REPORT_DATE = date(2026, 6, 30)

SECTORS = ["Banking & Finance", "Healthcare", "Logistics", "Technology", "Manufacturing", "Retail & F&B"]
ROLES = {
    "Banking & Finance": [("Relationship Manager", "Permanent", 6500), ("Compliance Analyst", "Permanent", 5800),
                          ("Accounts Executive", "Contract", 3800)],
    "Healthcare": [("Staff Nurse", "Permanent", 4200), ("Patient Service Associate", "Contract", 2600),
                   ("Clinic Assistant", "Temp", 2300)],
    "Logistics": [("Warehouse Supervisor", "Permanent", 3600), ("Logistics Coordinator", "Contract", 3100),
                  ("Forklift Operator", "Temp", 2400)],
    "Technology": [("Software Engineer", "Permanent", 7200), ("IT Support Engineer", "Contract", 4000),
                   ("Data Analyst", "Permanent", 5600)],
    "Manufacturing": [("Production Technician", "Contract", 2900), ("Quality Engineer", "Permanent", 4800),
                      ("Maintenance Technician", "Permanent", 3400)],
    "Retail & F&B": [("Retail Associate", "Temp", 2100), ("Store Manager", "Permanent", 4000),
                     ("Customer Service Executive", "Contract", 2700)],
}
SOURCES = ["LinkedIn", "JobStreet", "MyCareersFuture", "Referral", "BrightPath Website", "Job Fair", "Talent Database"]
SOURCE_W = [25, 20, 15, 10, 12, 6, 12]
SOURCE_QUALITY = {"LinkedIn": 1.0, "JobStreet": 0.85, "MyCareersFuture": 0.8, "Referral": 1.6,
                  "BrightPath Website": 1.05, "Job Fair": 0.7, "Talent Database": 1.25}
SOURCE_COST = {"LinkedIn": 38, "JobStreet": 22, "MyCareersFuture": 6, "Referral": 55,
               "BrightPath Website": 9, "Job Fair": 31, "Talent Database": 12}
RECRUITERS = [
    ("R01", "Aisha Rahman", "Finance & Tech Desk", 6),
    ("R02", "Daniel Tan", "Finance & Tech Desk", 5),
    ("R03", "Priya Nair", "Healthcare Desk", 5),
    ("R04", "Marcus Lim", "Industrial Desk", 6),
    ("R05", "Siti Hamid", "Industrial Desk", 5),
    ("R06", "Kevin Ong", "Retail Desk", 4),
]
RECRUITER_SECTORS = {
    "R01": ["Banking & Finance", "Technology"], "R02": ["Technology", "Banking & Finance"],
    "R03": ["Healthcare"], "R04": ["Logistics", "Manufacturing"], "R05": ["Manufacturing", "Logistics"],
    "R06": ["Retail & F&B"],
}
CLIENTS = {
    "Banking & Finance": ["Kestrel Wealth", "Orchard Crest Bank", "Silverline Insurance"],
    "Healthcare": ["Azurepoint Health", "Lumora Clinics", "Meridian Care Group"],
    "Logistics": ["Tanjong Freight", "Corallis Logistics", "Bayfront Cold Chain"],
    "Technology": ["Quillon Systems", "Nimbusly Cloud", "Arcwave Analytics"],
    "Manufacturing": ["Pinnacrest Manufacturing", "Veritek Precision", "Solano Polymers"],
    "Retail & F&B": ["Halcyon Retail", "Verdane Foods", "Kopi Lane Cafes"],
}

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")


def _recruiter_for(rng, sector):
    ids = [r for r, secs in RECRUITER_SECTORS.items() if sector in secs]
    return rng.choice(ids)


def _rand_date(rng, start, end):
    return start + timedelta(days=rng.randint(0, (end - start).days))


# --------------------------------------------------------------------------- workbook helpers
def _write_rows(ws, headers, rows, table_name=None, formats=None, widths=None, start_row=1):
    formats = formats or {}
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=c, value=h)
        if table_name is None:
            cell.font = Font(bold=True)
    for r, row in enumerate(rows, start_row + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            fmt = formats.get(headers[c - 1])
            if fmt:
                cell.number_format = fmt
            elif isinstance(v, date):
                cell.number_format = "yyyy-mm-dd"
    last = f"{get_column_letter(len(headers))}{start_row + len(rows)}"
    if table_name:
        t = Table(displayName=table_name, ref=f"A{start_row}:{last}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(t)
    for c, h in enumerate(headers, 1):
        w = (widths or {}).get(h) or max(11, min(34, len(h) + 3))
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


def _about_sheet(wb, title, purpose, dictionary):
    ws = wb.create_sheet("About_This_Data", 0)
    ws["A1"] = f"{COMPANY} — {title}"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = FICTION_NOTE
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A2:C2")
    ws.row_dimensions[2].height = 48
    ws["A3"] = purpose
    ws["A3"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A3:C3")
    ws.row_dimensions[3].height = 48
    ws["A5"], ws["B5"], ws["C5"] = "Sheet / Column", "Type", "Meaning"
    for c in ("A5", "B5", "C5"):
        ws[c].fill = HEADER_FILL
        ws[c].font = HEADER_FONT
    for i, (col, typ, meaning) in enumerate(dictionary, 6):
        ws.cell(row=i, column=1, value=col)
        ws.cell(row=i, column=2, value=typ)
        ws.cell(row=i, column=3, value=meaning).alignment = Alignment(wrap_text=True)
        if len(meaning) > 85:  # Excel does not auto-fit wrapped rows on open
            ws.row_dimensions[i].height = 15 * -(-len(meaning) // 85)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 90


def _prompt_log_sheet(wb):
    ws = wb.create_sheet("Prompt_Log")
    headers = ["Step", "Mode (Chat only / Plan / Allow editing)", "Prompt", "What Copilot did", "How I verified", "Decision (Keep / Revise / Undo)"]
    _write_rows(ws, headers, [], None, widths={"Prompt": 60, "What Copilot did": 40, "How I verified": 40})
    for c in range(1, len(headers) + 1):
        ws.cell(row=1, column=c).fill = HEADER_FILL
        ws.cell(row=1, column=c).font = HEADER_FONT


def _save(wb, folder, name):
    out = os.path.join(folder, "data")
    os.makedirs(out, exist_ok=True)
    if "Sheet" in wb.sheetnames and wb["Sheet"].max_row == 1 and wb["Sheet"]["A1"].value is None:
        del wb["Sheet"]
    path = os.path.join(out, name)
    wb.save(path)
    return path


# =========================================================================== LAB 1
def _lab1():
    rng = random.Random(3401)
    rows = []
    start, end = date(2026, 1, 5), date(2026, 6, 26)
    for i in range(1, 241):
        sector = rng.choices(SECTORS, [16, 15, 17, 20, 14, 18])[0]
        title, etype, base = rng.choice(ROLES[sector])
        source = rng.choices(SOURCES, SOURCE_W)[0]
        rec = _recruiter_for(rng, sector)
        applied = _rand_date(rng, start, end)
        yrs = max(0, round(rng.gauss(5, 3)))
        exp_sal = int(round(base * (0.85 + yrs * 0.03) * rng.uniform(0.92, 1.12), -1))
        q = SOURCE_QUALITY[source]
        age = (REPORT_DATE - applied).days
        r = rng.random()
        if age < 10 and r < 0.6:
            stage = "Applied"
        elif rng.random() < 0.42 / q:
            stage = "Rejected"
        elif rng.random() < 0.07:
            stage = "Withdrawn"
        else:
            s = rng.random() * q
            stage = "Screening" if s < 0.25 else "Interview" if s < 0.62 else "Offer" if s < 0.80 else "Hired"
        score = None
        if stage in ("Interview", "Offer", "Hired"):
            score = {"Interview": rng.choice([2, 3, 3, 4]), "Offer": rng.choice([3, 4, 4, 5]),
                     "Hired": rng.choice([4, 4, 5, 5])}[stage]
        elif stage == "Rejected" and rng.random() < 0.3:
            score = rng.choice([1, 2, 2, 3])
        days = min(age, max(1, int(rng.gauss({"Applied": 4, "Screening": 9, "Interview": 18, "Offer": 26,
                                                 "Hired": 33, "Rejected": 12, "Withdrawn": 15}[stage], 5))))
        rows.append([f"APP-{i:04d}", applied, f"JOB-{1000 + ROLES[sector].index((title, etype, base)) + 10 * SECTORS.index(sector)}",
                     title, sector, etype, source, dict((r[0], r[1]) for r in RECRUITERS)[rec], yrs, exp_sal,
                     stage, score, days])
    rows.sort(key=lambda r: (r[1], r[0]))
    # renumber in date order so IDs read naturally
    for i, r in enumerate(rows, 1):
        r[0] = f"APP-{i:04d}"
    # deliberate data-quality issues for the lab to find (no row is deleted)
    variant_rows = [r for r in rows if r[6] == "LinkedIn"][3:6]
    for r in variant_rows:
        r[6] = "Linked In"
    js_rows = [r for r in rows if r[6] == "JobStreet"][4:6]
    for r in js_rows:
        r[6] = "JobStreet "
    dup_src = [r for r in rows if r[10] == "Rejected"][7:9]
    dups = [list(r) for r in dup_src]
    rows = rows + dups
    headers = ["Application_ID", "Applied_Date", "Job_ID", "Job_Title", "Client_Sector", "Employment_Type", "Source",
               "Recruiter", "Years_Experience", "Expected_Salary", "Stage", "Interview_Score", "Days_in_Pipeline"]

    def clean(s):
        return s.strip().replace("Linked In", "LinkedIn")
    uniq = {}
    for r in rows:
        uniq.setdefault(r[0], r)
    ur = list(uniq.values())
    by_source = {}
    for r in ur:
        by_source[clean(r[6])] = by_source.get(clean(r[6]), 0) + 1
    hired = [r for r in ur if r[10] == "Hired"]
    hires_by_source = {}
    for r in hired:
        hires_by_source[clean(r[6])] = hires_by_source.get(clean(r[6]), 0) + 1
    by_sector = {}
    for r in ur:
        by_sector[r[4]] = by_sector.get(r[4], 0) + 1
    tech_interview = sum(1 for r in ur if r[4] == "Technology" and r[10] == "Interview")
    exp = dict(
        rows=len(rows), unique=len(ur), dup_ids=sorted({r[0] for r in dups}),
        hired=len(hired), linkedin_clean=by_source["LinkedIn"], linkedin_raw_exact=sum(1 for r in ur if r[6] == "LinkedIn"),
        variants_linkedin=len(variant_rows), variants_jobstreet=len(js_rows),
        top_sector=max(by_sector, key=by_sector.get), top_sector_n=max(by_sector.values()),
        top_hire_source=max(hires_by_source, key=hires_by_source.get), top_hire_source_n=max(hires_by_source.values()),
        tech_interview=tech_interview, first_date=min(r[1] for r in rows), last_date=max(r[1] for r in rows),
        score5=sum(1 for r in ur if r[11] == 5),
        linkedin_clean_all=sum(1 for r in rows if clean(r[6]) == "LinkedIn"),
        jobstreet_clean_all=sum(1 for r in rows if clean(r[6]) == "JobStreet"),
        tech_interview_all=sum(1 for r in rows if r[4] == "Technology" and r[10] == "Interview"),
        hired_all=sum(1 for r in rows if r[10] == "Hired"),
        dup_detail=[(r[0], r[6], r[10]) for r in dups],
    )
    return headers, rows, exp


def build_lab1(folder):
    headers, rows, _ = _lab1()
    wb = Workbook()
    ws = wb.active
    ws.title = "Applications"
    _write_rows(ws, headers, rows, None, {"Expected_Salary": '"S$"#,##0'}, {"Job_Title": 26, "Client_Sector": 18})
    _about_sheet(wb, "Candidate Pipeline (Jan–Jun 2026)",
                 "Purpose: one row per candidate application received by BrightPath's permanent and flexible "
                 "staffing desks. Use it in Lab 1 to explore, question, highlight, sort and filter pipeline data with Copilot. "
                 "The data is deliberately NOT yet formatted as an Excel table, and contains a few quality issues to find.",
                 [("Application_ID", "Text", "Unique application key (APP-0001 …). Should not repeat."),
                  ("Applied_Date", "Date", "Date the application was received."),
                  ("Job_ID / Job_Title", "Text", "Job order the candidate applied for."),
                  ("Client_Sector", "Text", "Sector of the hiring client."),
                  ("Employment_Type", "Text", "Permanent, Contract or Temp."),
                  ("Source", "Text", "Channel the application came from. Canonical values: " + ", ".join(SOURCES) + "."),
                  ("Recruiter", "Text", "BrightPath recruiter who owns the application."),
                  ("Years_Experience", "Number", "Candidate's years of relevant experience."),
                  ("Expected_Salary", "S$ / month", "Candidate's expected monthly base salary."),
                  ("Stage", "Text", "Current stage: Applied, Screening, Interview, Offer, Hired, Rejected, Withdrawn."),
                  ("Interview_Score", "1–5", "Panel score; blank if the candidate was not interviewed."),
                  ("Days_in_Pipeline", "Days", "Days since application to current stage or close.")])
    _prompt_log_sheet(wb)
    return [_save(wb, folder, "BrightPath_Candidate_Pipeline.xlsx")]


# =========================================================================== LAB 2
def _lab2():
    rng = random.Random(3402)
    reqs = []
    for i in range(1, 61):
        sector = rng.choices(SECTORS, [17, 15, 16, 20, 15, 17])[0]
        title, _etype, base = rng.choice(ROLES[sector])
        client = rng.choice(CLIENTS[sector])
        rec = _recruiter_for(rng, sector)
        openings = rng.choices([1, 1, 1, 2, 2, 3, 4, 5], k=1)[0]
        open_d = _rand_date(rng, date(2026, 1, 2), date(2026, 6, 15))
        budget = int(round(base * rng.uniform(0.95, 1.25), -2))
        fee = rng.choice([0.18, 0.20, 0.20, 0.22, 0.25])
        r = rng.random()
        age = (REPORT_DATE - open_d).days
        if age > 25 and r < 0.55:
            status = "Filled"
        elif r < 0.62:
            status = "Cancelled"
        elif r < 0.70:
            status = "On Hold"
        else:
            status = "Open"
        if status == "Filled":
            filled = openings
            close = open_d + timedelta(days=max(9, min(age, int(rng.gauss(34 if sector != "Technology" else 46, 10)))))
        elif status == "Cancelled":
            filled = rng.choice([0, 0, min(1, openings - 1)])
            close = open_d + timedelta(days=max(5, min(age, rng.randint(10, 50))))
        else:
            filled = rng.randint(0, max(0, openings - 1))
            close = None
        reqs.append([f"JOB-{2000 + i}", client, sector, title, rec, openings, filled, open_d, close, status, budget, fee])
    reqs.sort(key=lambda r: r[7])
    for i, r in enumerate(reqs, 1):
        r[0] = f"JOB-{2000 + i}"
    headers = ["Job_ID", "Client", "Sector", "Job_Title", "Recruiter_ID", "Openings", "Filled", "Open_Date",
               "Close_Date", "Status", "Monthly_Salary_Budget", "Fee_Rate"]

    def days_open(r):
        return ((r[8] or REPORT_DATE) - r[7]).days
    filled_rows = [r for r in reqs if r[9] == "Filled"]
    fee_total = sum(r[6] * r[10] * 12 * r[11] for r in reqs)
    overdue = [r for r in reqs if r[9] == "Open" and days_open(r) > 45]
    names = {r[0]: r[1] for r in RECRUITERS}
    fee_by_rec = {}
    for r in reqs:
        fee_by_rec[names[r[4]]] = fee_by_rec.get(names[r[4]], 0) + r[6] * r[10] * 12 * r[11]
    sample = next(r for r in reqs if r[9] == "Filled" and r[6] >= 2)
    exp = dict(
        rows=len(reqs), filled_n=len(filled_rows),
        avg_days_filled=round(statistics.mean(days_open(r) for r in filled_rows), 1),
        fee_total=round(fee_total, 2), overdue_n=len(overdue), overdue_ids=[r[0] for r in overdue],
        top_rec=max(fee_by_rec, key=fee_by_rec.get), top_rec_fee=round(max(fee_by_rec.values()), 2),
        openings=sum(r[5] for r in reqs), filled_total=sum(r[6] for r in reqs),
        sample_id=sample[0], sample_fee=round(sample[6] * sample[10] * 12 * sample[11], 2),
        sample_days=days_open(sample), sample_detail=(sample[6], sample[10], sample[11]),
        status_counts={s: sum(1 for r in reqs if r[9] == s) for s in ("Open", "Filled", "On Hold", "Cancelled")},
    )
    return headers, reqs, exp


def build_lab2(folder):
    headers, rows, _ = _lab2()
    wb = Workbook()
    ws = wb.active
    ws.title = "Requisitions"
    _write_rows(ws, headers, rows, "tblRequisitions",
                {"Monthly_Salary_Budget": '"S$"#,##0', "Fee_Rate": "0%"}, {"Client": 24, "Job_Title": 26})
    ws2 = wb.create_sheet("Recruiters")
    _write_rows(ws2, ["Recruiter_ID", "Recruiter_Name", "Desk", "Monthly_Target_Placements"],
                [list(r) for r in RECRUITERS], "tblRecruiters", widths={"Desk": 22})
    ws3 = wb.create_sheet("Control")
    ws3["A1"], ws3["B1"] = "Report_Date", REPORT_DATE
    ws3["B1"].number_format = "yyyy-mm-dd"
    ws3["A1"].font = Font(bold=True)
    ws3["D1"] = "Use this sheet for your independent check formulas (Lab 2, Step 6)."
    ws3.column_dimensions["A"].width = 34
    ws3.column_dimensions["B"].width = 16
    _about_sheet(wb, "Permanent-Recruitment Job Requisitions (H1 2026)",
                 "Purpose: one row per client job order (requisition) on BrightPath's permanent-recruitment desk. BrightPath "
                 "charges a placement fee equal to Fee_Rate x first-year base salary (Monthly_Salary_Budget x 12) for every "
                 "position filled. Use it in Lab 2 to generate and verify formula columns and lookups with Copilot.",
                 [("Job_ID", "Text", "Unique requisition key."), ("Client / Sector", "Text", "Fictional hiring client and its sector."),
                  ("Recruiter_ID", "Text", "Owner; look up the name in tblRecruiters."),
                  ("Openings / Filled", "Count", "Positions requested and positions filled so far."),
                  ("Open_Date / Close_Date", "Date", "Close_Date is blank while the requisition is Open or On Hold."),
                  ("Status", "Text", "Open, Filled, On Hold or Cancelled."),
                  ("Monthly_Salary_Budget", "S$ / month", "Client's budgeted monthly base salary per position."),
                  ("Fee_Rate", "%", "Placement fee as a % of first-year base salary."),
                  ("Control!B1 Report_Date", "Date", "Fixed reporting date 2026-06-30 used for ageing open requisitions.")])
    _prompt_log_sheet(wb)
    return [_save(wb, folder, "BrightPath_Job_Requisitions.xlsx")]


# =========================================================================== LAB 3
def _months(y0, m0, n):
    out = []
    y, m = y0, m0
    for _ in range(n):
        out.append(date(y, m, 1))
        m += 1
        if m == 13:
            y, m = y + 1, 1
    return out


def _split_total(weights, total):
    """Whole numbers in proportion to weights that add up exactly to total (largest remainder)."""
    raw = [total * w / sum(weights) for w in weights]
    out = [int(x) for x in raw]
    for i in sorted(range(len(raw)), key=lambda i: raw[i] - out[i], reverse=True)[:total - sum(out)]:
        out[i] += 1
    return out


def _lab3():
    rng = random.Random(3403)
    months = _months(2025, 1, 18)
    base = {"Banking & Finance": 120, "Healthcare": 150, "Logistics": 170, "Technology": 210,
            "Manufacturing": 130, "Retail & F&B": 190}
    rows = []
    for mi, m in enumerate(months):
        season = 1 + 0.12 * (m.month in (1, 3, 9)) - 0.10 * (m.month in (6, 12))
        for s in SECTORS:
            growth = 1 + (0.025 * mi if s in ("Healthcare", "Technology") else 0.004 * mi)
            if s == "Retail & F&B" and m.year == 2026:
                growth *= 0.82
            apps = int(base[s] * season * growth * rng.uniform(0.9, 1.1))
            screened = int(apps * rng.uniform(0.45, 0.55))
            interviews = int(screened * rng.uniform(0.38, 0.48))
            offers = int(interviews * rng.uniform(0.32, 0.42))
            hire_ratio = rng.uniform(0.70, 0.86) if s != "Technology" else rng.uniform(0.58, 0.70)
            hires = max(1, int(round(offers * hire_ratio)))
            tth = round(rng.gauss({"Technology": 41, "Banking & Finance": 37, "Healthcare": 33}.get(s, 24), 3), 1)
            cph = int(round(rng.gauss({"Technology": 2400, "Banking & Finance": 2100}.get(s, 1100), 120), -1))
            rows.append([m, m.strftime("%Y-%m"), s, apps, screened, interviews, offers, hires, tth, cph])
    headers = ["Month", "Month_Label", "Sector", "Applications", "Screened", "Interviews", "Offers", "Hires",
               "Avg_Time_to_Hire_Days", "Cost_per_Hire"]
    h1 = [r for r in rows if r[0].year == 2026]
    # Jan-Jun 2026 by source: applications and hires add up to tblMonthly's 2026 totals,
    # so a learner who reconciles the two tables finds they agree
    rng2 = random.Random(34031)
    apps_w = [SOURCE_W[i] * rng2.uniform(0.9, 1.1) for i in range(len(SOURCES))]
    src_apps = _split_total(apps_w, sum(r[3] for r in h1))
    hires_w = [a * SOURCE_QUALITY[s] * rng2.uniform(0.9, 1.1) for s, a in zip(SOURCES, src_apps)]
    src_hires = _split_total(hires_w, sum(r[7] for r in h1))
    src_rows = [[s, a, h, int(round(a * SOURCE_COST[s] * rng2.uniform(0.95, 1.05), -2))]
                for s, a, h in zip(SOURCES, src_apps, src_hires)]
    by_sector = {}
    for r in h1:
        by_sector[r[2]] = by_sector.get(r[2], 0) + r[7]
    by_month = {}
    for r in rows:
        by_month[r[1]] = by_month.get(r[1], 0) + r[7]
    funnel = [sum(r[i] for r in h1) for i in (3, 4, 5, 6, 7)]
    hire_rate = {r[0]: r[2] / r[1] for r in src_rows}
    cps = {r[0]: r[3] / r[2] for r in src_rows}
    exp = dict(rows=len(rows), h1_hires=sum(r[7] for r in h1), top_sector=max(by_sector, key=by_sector.get),
               top_sector_hires=max(by_sector.values()), by_sector=by_sector, peak_month=max(by_month, key=by_month.get),
               peak_month_hires=max(by_month.values()), funnel=funnel,
               best_rate_source=max(hire_rate, key=hire_rate.get), best_rate=round(max(hire_rate.values()) * 100, 1),
               lowest_cps_source=min(cps, key=cps.get), lowest_cps=round(min(cps.values())),
               total_2025=sum(r[7] for r in rows if r[0].year == 2025))
    return headers, rows, src_rows, exp


def build_lab3(folder):
    headers, rows, src_rows, _ = _lab3()
    wb = Workbook()
    ws = wb.active
    ws.title = "Monthly_Metrics"
    _write_rows(ws, headers, rows, "tblMonthly", {"Month": "mmm-yyyy", "Cost_per_Hire": '"S$"#,##0',
                                                  "Avg_Time_to_Hire_Days": "0.0"}, {"Sector": 18})
    ws2 = wb.create_sheet("Source_Summary")
    _write_rows(ws2, ["Source", "Applications", "Hires", "Sourcing_Spend"], src_rows, "tblSources",
                {"Sourcing_Spend": '"S$"#,##0'}, {"Source": 20})
    wb.create_sheet("Charts")
    _about_sheet(wb, "Monthly Recruitment Metrics (Jan 2025 – Jun 2026)",
                 "Purpose: one row per month per client sector with funnel counts and efficiency measures, plus a "
                 "Jan–Jun 2026 sourcing-channel summary. Use it in Lab 3 to create and refine charts with Copilot. "
                 "The Charts sheet starts empty: create each chart there with Copilot.",
                 [("Month / Month_Label", "Date / Text", "First day of the month and its yyyy-mm label."),
                  ("Sector", "Text", "Client sector."),
                  ("Applications → Hires", "Count", "Funnel counts for the month: Applications, Screened, Interviews, Offers, Hires."),
                  ("Avg_Time_to_Hire_Days", "Days", "Average days from application to accepted offer for hires that month."),
                  ("Cost_per_Hire", "S$", "Average sourcing + recruiter cost per hire."),
                  ("Source_Summary sheet", "Table", "Applications, hires and sourcing spend by channel, Jan–Jun 2026.")])
    _prompt_log_sheet(wb)
    return [_save(wb, folder, "BrightPath_Recruitment_Metrics.xlsx")]


# =========================================================================== LAB 4 (workbook rules)
REGIONS = ["CENTRAL", "EAST", "NORTH", "NORTH-EAST", "WEST"]
REGION_VARIANTS = {
    "CENTRAL": ["Central", "central", "CENTRAL", "C", "Central Region", "Central "],
    "EAST": ["East", "east", "E", "East Region"],
    "NORTH": ["North", "north", "N", "North "],
    "NORTH-EAST": ["North-East", "Northeast", "North East", "NE", "N-E"],
    "WEST": ["West", "west", "W", "West Region"],
}
STATUSES = ["Open", "In Progress", "On Hold", "Completed", "Cancelled"]
STATUS_VARIANTS = {
    "Open": ["Open", "open", "Not started", "New"],
    "In Progress": ["In Progress", "in progress", "In-Progress", "WIP", "Ongoing"],
    "On Hold": ["On Hold", "on hold", "Paused", "Hold"],
    "Completed": ["Completed", "Complete", "Done", "completed", "Closed - filled"],
    "Cancelled": ["Cancelled", "Canceled", "cancelled", "Client cancelled"],
}
SERVICES = ["Permanent Search", "Contract Staffing", "Payroll Outsourcing", "RPO"]
LEGAL_SUFFIXES = [" Pte Ltd", " Pte. Ltd.", " PTE LTD", " Private Limited", "", ""]
RATE_CARD = [("Aisha Rahman", "Senior Consultant", 120), ("Daniel Tan", "Consultant", 95),
             ("Priya Nair", "Senior Consultant", 115), ("Marcus Lim", "Principal Consultant", 140),
             ("Siti Hamid", "Consultant", 90), ("Kevin Ong", "Associate Consultant", 75)]
_MONTHS3 = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
PAUSE_NOTE = "Client paused hiring pending headcount approval"


def _lab4():
    rng = random.Random(3414)
    clients = [c for s in SECTORS for c in CLIENTS[s][:2]]          # 12 clients
    rate = {o: r for o, _role, r in RATE_CARD}
    canon, raw = [], []
    for i in range(1, 65):
        client = rng.choice(clients)
        owner = rng.choice(list(rate))
        status = rng.choices(STATUSES, [12, 22, 6, 18, 6])[0]
        region = rng.choices(REGIONS, [34, 18, 14, 16, 18])[0]
        start = _rand_date(rng, date(2025, 10, 1), date(2026, 6, 15))
        budget = rng.choice(range(6000, 90001, 500))
        r = rate[owner]
        f = {"Completed": rng.uniform(0.75, 1.25), "In Progress": rng.uniform(0.2, 1.1),
             "On Hold": rng.uniform(0.1, 0.5), "Open": 0.0, "Cancelled": rng.uniform(0.0, 0.3)}[status]
        hours = int(round(budget / r * f)) if status != "Open" else rng.randint(0, 4)
        canon.append(dict(id=f"ORD-{300 + i}", client=client, service=rng.choice(SERVICES), owner=owner,
                          region=region, status=status, start=start, budget=budget, hours=hours,
                          value=hours * r))
    # Two On Hold orders are exported as "Waiting to start" with a note that says the work is paused.
    paused = [c for c in canon if c["status"] == "On Hold"][:2]
    plain_notes = ["Weekly update sent", "Shortlist due Friday", "Client feedback received", "Contract signed",
                   "Interviews scheduled", "Awaiting job description", "", "", "", ""]
    for c in canon:
        name = c["client"] + rng.choice(LEGAL_SUFFIXES)
        roll = rng.random()
        if roll < 0.12:
            name = name.upper()
        elif roll < 0.22:
            name = name.lower()
        if rng.random() < 0.18:
            name = name + " "
        own = c["owner"]
        roll = rng.random()
        if roll < 0.10:
            own = own + " "
        elif roll < 0.16:
            own = own.lower()
        elif roll < 0.20:
            own = own.replace(" ", "  ", 1)
        if c in paused:
            st, note = "Waiting to start", PAUSE_NOTE
        else:
            st, note = rng.choice(STATUS_VARIANTS[c["status"]]), rng.choice(plain_notes)
        d = c["start"]
        fmt = rng.choice(["date", "dmy", "dmonY", "iso", "dot"])
        sd = {"date": d, "dmy": d.strftime("%d/%m/%Y"), "dmonY": f"{d.day} {_MONTHS3[d.month - 1]} {d.year}",
              "iso": d.isoformat(), "dot": d.strftime("%d.%m.%Y")}[fmt]
        b = c["budget"]
        bfmt = rng.choice(["num", "sd", "sgd", "k", "dollar"])
        bud = {"num": b, "sd": f"S${b:,}", "sgd": f"SGD {b:,}", "k": f"{b / 1000:g}k", "dollar": f"${b:,}"}[bfmt]
        raw.append([c["id"], name, c["service"], own, rng.choice(REGION_VARIANTS[c["region"]]), st, sd, bud,
                    c["hours"], note])
    headers = ["Order_ID", "Client", "Service", "Owner", "Region", "Status", "Start_Date", "Budget", "Hours_Logged",
               "Notes"]
    over = [c for c in canon if c["value"] > c["budget"]]
    exp = dict(
        rows=len(canon), clients=len({c["client"] for c in canon}), owners=len({c["owner"] for c in canon}),
        status_counts={s: sum(1 for c in canon if c["status"] == s) for s in STATUSES},
        region_counts={g: sum(1 for c in canon if c["region"] == g) for g in REGIONS},
        total_budget=sum(c["budget"] for c in canon), total_value=sum(c["value"] for c in canon),
        over_n=len(over), over_ids=[c["id"] for c in over],
        paused_ids=[c["id"] for c in paused],
        raw_status_distinct=len({r[5] for r in raw}), raw_region_distinct=len({r[4] for r in raw}),
        raw_client_distinct=len({r[1] for r in raw}),
        top_client=max({c["client"] for c in canon}, key=lambda k: sum(x["budget"] for x in canon if x["client"] == k)),
    )
    exp["top_client_budget"] = sum(c["budget"] for c in canon if c["client"] == exp["top_client"])
    s = next(c for c in canon if c["status"] == "Completed" and c["value"] > c["budget"])
    exp["sample"] = dict(id=s["id"], owner=s["owner"], hours=s["hours"], rate=rate[s["owner"]], value=s["value"],
                         budget=s["budget"])
    return headers, raw, canon, exp


def build_lab4(folder):
    headers, raw, _canon, _ = _lab4()
    wb = Workbook()
    ws = wb.active
    ws.title = "Orders_Raw"
    _write_rows(ws, headers, raw, None, None, {"Client": 32, "Service": 20, "Owner": 16, "Notes": 44})
    ws2 = wb.create_sheet("Rate_Card")
    _write_rows(ws2, ["Owner", "Role", "Hourly_Rate"], RATE_CARD, "tblRateCard", {"Hourly_Rate": '"S$"#,##0'},
                {"Owner": 16, "Role": 22})
    _about_sheet(wb, "Client Orders - Raw Export (Oct 2025 - Jun 2026)",
                 "Purpose: one row per client order, exported from BrightPath's order system. Several people typed "
                 "these values, so the export is messy: client names, owners, regions, statuses, dates and budgets are "
                 "written in different ways. Use it in Lab 4 to clean the data with Copilot's automated cleanup and a "
                 ".Rules worksheet. Keep Orders_Raw unchanged - build clean copies on new sheets.",
                 [("Order_ID", "Text", "Unique order key (ORD-301 ...)."),
                  ("Client", "Text", "Client trading name, sometimes with a legal suffix such as Pte Ltd."),
                  ("Service", "Text", "Permanent Search, Contract Staffing, Payroll Outsourcing or RPO."),
                  ("Owner", "Text", "BrightPath consultant who owns the order. Matches Rate_Card[Owner] once cleaned."),
                  ("Region", "Text", "Singapore region of the client site. Valid values: " + ", ".join(REGIONS) + "."),
                  ("Status", "Text", "Valid values: " + ", ".join(STATUSES) + "."),
                  ("Start_Date", "Date or text", "Order start date. Text dates are day-first (dd/mm/yyyy)."),
                  ("Budget", "S$ or text", "Order budget in Singapore dollars, typed in several styles (e.g. 12.5k)."),
                  ("Hours_Logged", "Hours", "Consultant hours booked to the order so far."),
                  ("Notes", "Text", "Free-text notes. A note can change the meaning of the Status."),
                  ("Rate_Card sheet", "Table", "tblRateCard: hourly charge-out rate for each consultant.")])
    _prompt_log_sheet(wb)
    return [_save(wb, folder, "BrightPath_Client_Orders_Raw.xlsx")]


# =========================================================================== LAB 5 (built-in skills)
PLANS = [("Payroll Only", 1800), ("Payroll + HR Admin", 3600), ("Full HR Outsourcing", 7400)]
CHANNELS = ["Referral", "Accounting Partners", "LinkedIn Ads", "Events", "Website"]
CHANNEL_SPEND = {"Referral": 6500, "Accounting Partners": 19000, "LinkedIn Ads": 30000, "Events": 23000, "Website": 11000}
CHANNEL_W = [22, 26, 22, 14, 16]


def _lab5():
    rng = random.Random(3415)
    months = _months(2025, 1, 18)
    clients = []
    for i in range(1, 61):                    # clients already active on 1 Jan 2025
        plan, fee = rng.choices(PLANS, [45, 38, 17])[0]
        clients.append(dict(sign=_rand_date(rng, date(2023, 1, 1), date(2024, 12, 20)), plan=plan,
                            channel=rng.choices(CHANNELS, CHANNEL_W)[0], fee=fee, churn=None))
    spend_rows, month_rows = [], []
    for m0 in months:
        y, m = m0.year, m0.month
        m1 = date(y + (m == 12), m % 12 + 1, 1) - timedelta(days=1)
        active_start = [c for c in clients if c["sign"] < m0 and (c["churn"] is None)]
        churned = 0
        for c in active_start:
            p = {"Payroll Only": 0.024, "Payroll + HR Admin": 0.014, "Full HR Outsourcing": 0.008}[c["plan"]]
            if rng.random() < p:
                c["churn"] = _rand_date(rng, m0, m1)
                churned += 1
        new = rng.randint(3, 7) + (1 if y == 2026 else 0)
        for _ in range(new):
            plan, fee = rng.choices(PLANS, [45, 38, 17])[0]
            clients.append(dict(sign=_rand_date(rng, m0, m1), plan=plan, channel=rng.choices(CHANNELS, CHANNEL_W)[0],
                                fee=fee, churn=None))
        spend = 0
        for ch in CHANNELS:
            v = int(round(CHANNEL_SPEND[ch] * rng.uniform(0.8, 1.25), -1))
            spend += v
            spend_rows.append([m0, ch, v])
        active_end = [c for c in clients if c["sign"] <= m1 and (c["churn"] is None or c["churn"] > m1)]
        mrr = sum(c["fee"] for c in active_end)
        gm = round(rng.uniform(0.41, 0.47), 3)
        month_rows.append([m0, f"{y}-{m:02d}", len(active_start), new, churned, len(active_end), mrr, gm, spend])
    # client table with IDs in sign-up order and a scaled fee per client
    clients.sort(key=lambda c: c["sign"])
    crow = []
    for i, c in enumerate(clients, 1):
        emp = {"Payroll Only": rng.randint(15, 60), "Payroll + HR Admin": rng.randint(30, 120),
               "Full HR Outsourcing": rng.randint(60, 250)}[c["plan"]]
        c["id"] = f"CL-{i:04d}"
        crow.append([c["id"], c["sign"], c["channel"], c["plan"], emp, c["fee"],
                     "Churned" if c["churn"] else "Active", c["churn"]])
    active_by_plan = {p: sum(1 for c in clients if c["plan"] == p and c["churn"] is None) for p, _fee in PLANS}
    h1 = [r for r in month_rows if r[0] >= date(2026, 1, 1)]
    by_ch = {}
    for ch in CHANNELS:
        sp = sum(r[2] for r in spend_rows if r[1] == ch and r[0] >= date(2026, 1, 1))
        nw = sum(1 for c in clients if c["channel"] == ch and c["sign"] >= date(2026, 1, 1))
        by_ch[ch] = dict(spend=sp, new=nw)
    end = month_rows[-1]
    exp = dict(
        months=len(month_rows), clients=len(clients), active_end=end[5], mrr_end=end[6], mrr_jan25=month_rows[0][6],
        new_all=sum(r[3] for r in month_rows), churned_all=sum(r[4] for r in month_rows),
        new_h1=sum(r[3] for r in h1), churned_h1=sum(r[4] for r in h1), spend_h1=sum(r[8] for r in h1),
        gm_min=round(min(r[7] for r in month_rows) * 100, 1), gm_max=round(max(r[7] for r in month_rows) * 100, 1),
        by_channel=by_ch, top_channel=max(by_ch, key=lambda c: by_ch[c]["spend"]),
        low_channel=min(by_ch, key=lambda c: by_ch[c]["spend"]),
        active_by_plan=active_by_plan, top_plan=max(active_by_plan, key=active_by_plan.get),
    )
    return month_rows, crow, spend_rows, exp


def build_lab5(folder):
    month_rows, crow, spend_rows, _ = _lab5()
    wb = Workbook()
    ws = wb.active
    ws.title = "Client_Months"
    _write_rows(ws, ["Month", "Month_Label", "Clients_Start", "New_Clients", "Churned_Clients", "Clients_End", "MRR",
                     "Gross_Margin_Pct", "Sales_Marketing_Spend"], month_rows, "tblClientMonths",
                {"Month": "mmm-yyyy", "MRR": '"S$"#,##0', "Gross_Margin_Pct": "0.0%", "Sales_Marketing_Spend": '"S$"#,##0'})
    ws2 = wb.create_sheet("Clients")
    _write_rows(ws2, ["Client_ID", "Sign_Date", "Channel", "Plan", "Employees_Served", "Monthly_Fee", "Status",
                      "Churn_Date"], crow, "tblClients", {"Monthly_Fee": '"S$"#,##0'}, {"Channel": 20, "Plan": 22})
    ws3 = wb.create_sheet("Marketing_Spend")
    _write_rows(ws3, ["Month", "Channel", "Spend"], spend_rows, "tblSpend", {"Month": "mmm-yyyy", "Spend": '"S$"#,##0'},
                {"Channel": 20})
    _about_sheet(wb, "HR Outsourcing Clients and Growth",
                 "Purpose: BrightPath's HR-outsourcing business sells monthly payroll and HR services to client "
                 "companies. Use it in Lab 5 to try Copilot's built-in visual skills: @chart-design to plot charts from "
                 "a prompt and @theme-design to give them one consistent look.",
                 [("tblClientMonths", "Table", "One row per month, Jan 2025 - Jun 2026: clients at start, new, churned, "
                   "clients at end, MRR (monthly recurring revenue), gross margin % and sales & marketing spend."),
                  ("tblClients", "Table", "One row per client: sign date, acquisition channel, plan, employees served, "
                   "monthly fee, Active or Churned, churn date."),
                  ("tblSpend", "Table", "Sales and marketing spend by month and channel."),
                  ("Definitions", "Rule", "H1 2026 = Jan-Jun 2026. MRR = the monthly fees of all clients active at the end "
                   "of the month. Active clients = clients with Status Active (Clients_End in the latest month). "
                   "Churned = clients lost in the month.")])
    _prompt_log_sheet(wb)
    return [_save(wb, folder, "BrightPath_HR_Outsourcing_Finance.xlsx")]


# =========================================================================== LAB 6 (custom skill: Pareto chart)
# 500 job postings in 16 job categories. The counts are fixed (and all different) so the Pareto order is unambiguous
# and the 80% point is easy to verify: the long tail is deliberate.
JOB_CATEGORIES = [
    ("Customer Service", 145), ("Software Engineering", 95), ("Administration", 65), ("Sales & Business Development", 50),
    ("Accounting & Finance", 35), ("Nursing & Healthcare", 25), ("Logistics & Warehouse", 19), ("IT Support", 15),
    ("Data & Analytics", 12), ("Human Resources", 10), ("Marketing", 8), ("Manufacturing Engineering", 7),
    ("Retail Operations", 6), ("Legal & Compliance", 4), ("Project Management", 3), ("Procurement", 1),
]
_ALL = list(SECTORS)
CATEGORY_SECTORS = {
    "Customer Service": ["Banking & Finance", "Retail & F&B", "Healthcare", "Technology", "Logistics"],
    "Software Engineering": ["Technology", "Technology", "Banking & Finance"],
    "Administration": _ALL,
    "Sales & Business Development": ["Technology", "Banking & Finance", "Retail & F&B", "Manufacturing"],
    "Accounting & Finance": ["Banking & Finance", "Banking & Finance", "Manufacturing", "Retail & F&B"],
    "Nursing & Healthcare": ["Healthcare"],
    "Logistics & Warehouse": ["Logistics", "Logistics", "Retail & F&B", "Manufacturing"],
    "IT Support": ["Technology", "Technology", "Banking & Finance", "Healthcare"],
    "Data & Analytics": ["Technology", "Banking & Finance"],
    "Human Resources": _ALL,
    "Marketing": ["Retail & F&B", "Technology"],
    "Manufacturing Engineering": ["Manufacturing"],
    "Retail Operations": ["Retail & F&B"],
    "Legal & Compliance": ["Banking & Finance", "Healthcare"],
    "Project Management": ["Technology", "Manufacturing", "Logistics"],
    "Procurement": ["Manufacturing", "Logistics"],
}


def _lab6():
    assert sum(n for _c, n in JOB_CATEGORIES) == 500
    for seed in range(3416, 3516):                  # first seed whose Client_Sector counts are all different, with a clear leader
        rng = random.Random(seed)
        cats = [c for c, n in JOB_CATEGORIES for _ in range(n)]
        rng.shuffle(cats)
        rows = []
        for cat in cats:
            rows.append([None, _rand_date(rng, date(2026, 1, 5), date(2026, 6, 26)), cat,
                         rng.choice(CATEGORY_SECTORS[cat]), rng.choices(["Permanent", "Contract", "Temp"], [55, 30, 15])[0],
                         rng.choices(["Filled", "Open", "Closed"], [52, 33, 15])[0],
                         rng.choices([1, 2, 3, 4, 5], [50, 25, 12, 8, 5])[0]])
        sec = {}
        for r in rows:
            sec[r[3]] = sec.get(r[3], 0) + 1
        counts = sorted(sec.values(), reverse=True)
        if len(set(counts)) == len(counts) and counts[0] - counts[1] >= 10:
            break
    rows.sort(key=lambda r: (r[1], r[2]))
    for i, r in enumerate(rows, 1):
        r[0] = f"JOB-{i:04d}"
    headers = ["Job_ID", "Posted_Date", "Job_Category", "Client_Sector", "Employment_Type", "Status", "Openings"]

    def pareto(counts, threshold=0.8):
        total = sum(n for _k, n in counts)
        cum, out, n_reach = 0, [], None
        for i, (k, n) in enumerate(counts, 1):
            cum += n
            out.append((k, n, round(n / total * 100, 1), round(cum / total * 100, 1)))
            if n_reach is None and cum / total >= threshold:
                n_reach = i
        return out, n_reach

    cat_counts = [(c, sum(1 for r in rows if r[2] == c)) for c, _n in JOB_CATEGORIES]
    cat_counts.sort(key=lambda kv: -kv[1])
    cat_table, n80 = pareto(cat_counts)
    sec_counts = sorted(sec.items(), key=lambda kv: -kv[1])
    sec_table, sec_n80 = pareto(sec_counts)
    small = [(k, n) for k, n, share, _c in cat_table if share < 2.0]
    exp = dict(
        rows=len(rows), categories=len(cat_counts), by_cat=cat_table, top_cat=cat_table[0][0], top_n=cat_table[0][1],
        top_share=cat_table[0][2], top3_share=cat_table[2][3], n80=n80, cum80=cat_table[n80 - 1][3],
        last80=cat_table[n80 - 1][0], vital=[c[0] for c in cat_table[:n80]],
        by_sector=sec_table, top_sector=sec_table[0][0], top_sector_n=sec_table[0][1], top_sector_share=sec_table[0][2],
        sector_n80=sec_n80, sector_cum80=sec_table[sec_n80 - 1][3], sectors=len(sec_table),
        small_n=len(small), small_total=sum(n for _k, n in small), small_share=round(sum(n for _k, n in small) / len(rows) * 100, 1),
        first_date=min(r[1] for r in rows), last_date=max(r[1] for r in rows),
    )
    return headers, rows, exp


def build_lab6(folder):
    headers, rows, _ = _lab6()
    wb = Workbook()
    ws = wb.active
    ws.title = "Job_Postings"
    _write_rows(ws, headers, rows, "tblJobs", None, {"Job_Category": 30, "Client_Sector": 20, "Employment_Type": 17})
    _about_sheet(wb, "Job Postings, Jan - Jun 2026",
                 "Purpose: one row per job posting BrightPath received from clients between January and June 2026 - 500 "
                 "postings in 16 job categories. Use it in Lab 6 to build a custom Copilot skill, pareto-chart, that plots a "
                 "Pareto chart of any category column, and to use that skill on Job_Category and Client_Sector.",
                 [("tblJobs", "Table", "One row per job posting."),
                  ("Job_ID", "Text", "Unique key (JOB-0001 ...)."),
                  ("Posted_Date", "Date", "When the client posted the job to BrightPath."),
                  ("Job_Category", "Text", "One of 16 job categories, e.g. Customer Service. The column you will plot."),
                  ("Client_Sector", "Text", "The hiring client's sector (6 values)."),
                  ("Employment_Type", "Text", "Permanent, Contract or Temp."),
                  ("Status", "Text", "Filled, Open or Closed."),
                  ("Openings", "Number", "Positions the client wants to fill (1 to 5)."),
                  ("Pareto chart", "Rule", "Columns = the count of each category, sorted high to low, a different colour "
                   "for each column. Line = cumulative % of the total on a secondary axis from 0% to 100%.")])
    _prompt_log_sheet(wb)
    # Open on the data: Copilot reads the active sheet first, so lead with tblJobs, not the notes.
    wb.move_sheet("Job_Postings", offset=-wb.sheetnames.index("Job_Postings"))
    wb.active = wb["Job_Postings"]
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == "Job_Postings"
    return [_save(wb, folder, "BrightPath_Job_Postings.xlsx")]


# =========================================================================== LAB 7 (interactive recruitment dashboard)
def _lab7():
    rng = random.Random(3405)
    rows = []
    rec_names = {r[0]: r[1] for r in RECRUITERS}
    for i in range(1, 601):
        s = rng.choices(SECTORS, [16, 15, 17, 20, 14, 18])[0]
        src = rng.choices(SOURCES, SOURCE_W)[0]
        rec = _recruiter_for(rng, s)
        applied = _rand_date(rng, date(2026, 1, 2), date(2026, 6, 20))
        q = SOURCE_QUALITY[src]
        stage = "Applied"
        if rng.random() < 0.55 * min(1.3, q):
            stage = "Screened"
            if rng.random() < 0.45 * min(1.3, q):
                stage = "Interviewed"
                if rng.random() < 0.42:
                    stage = "Offered"
        offer_made = "Yes" if stage == "Offered" else "No"
        accepted, hire_date, tth = "", None, None
        if offer_made == "Yes":
            p_acc = 0.62 if s == "Technology" else 0.80
            if rng.random() < p_acc:
                accepted = "Yes"
                tth = max(12, int(rng.gauss({"Technology": 40, "Banking & Finance": 36, "Healthcare": 31}.get(s, 24), 6)))
                hire_date = applied + timedelta(days=tth)
                if hire_date > REPORT_DATE:
                    accepted, hire_date, tth = "Pending", None, None
            else:
                accepted = "No"
        final = "Hired" if accepted == "Yes" else stage
        rows.append([f"A26-{i:04d}", applied, applied.strftime("%Y-%m"), s, src, rec_names[rec], final, offer_made,
                     accepted, hire_date, tth, SOURCE_COST[src]])
    rows.sort(key=lambda r: (r[1], r[0]))
    for i, r in enumerate(rows, 1):
        r[0] = f"A26-{i:04d}"
    headers = ["Application_ID", "Applied_Date", "Applied_Month", "Sector", "Source", "Recruiter", "Furthest_Stage",
               "Offer_Made", "Offer_Accepted", "Hire_Date", "Time_to_Hire_Days", "Sourcing_Cost"]
    hires = [r for r in rows if r[6] == "Hired"]
    offers = [r for r in rows if r[7] == "Yes"]
    decided = [r for r in offers if r[8] in ("Yes", "No")]
    by_month = {}
    for r in hires:
        k = r[9].strftime("%Y-%m")
        by_month[k] = by_month.get(k, 0) + 1
    by_sector = {}
    for r in hires:
        by_sector[r[3]] = by_sector.get(r[3], 0) + 1
    by_rec = {}
    for r in hires:
        by_rec[r[5]] = by_rec.get(r[5], 0) + 1
    tech_rows = [r for r in rows if r[3] == "Technology"]
    tech_hires = [r for r in tech_rows if r[6] == "Hired"]
    exp = dict(apps=len(rows), hires=len(hires), offers=len(offers), decided=len(decided),
               accepted=sum(1 for r in offers if r[8] == "Yes"), pending=sum(1 for r in offers if r[8] == "Pending"),
               oar=round(sum(1 for r in offers if r[8] == "Yes") / len(decided) * 100, 1),
               avg_tth=round(statistics.mean(r[10] for r in hires), 1),
               total_cost=sum(r[11] for r in rows), cph=round(sum(r[11] for r in rows) / len(hires), 2),
               by_month=dict(sorted(by_month.items())), top_sector=max(by_sector, key=by_sector.get),
               by_sector=by_sector, top_rec=max(by_rec, key=by_rec.get), top_rec_n=max(by_rec.values()),
               tech_apps=len(tech_rows), tech_hires=len(tech_hires),
               tech_avg_tth=round(statistics.mean(r[10] for r in tech_hires), 1) if tech_hires else 0)
    targets = [[m, t] for m, t in zip(["2026-01", "2026-02", "2026-03", "2026-04", "2026-05", "2026-06"],
                                      [0, 8, 9, 9, 9, 9])]
    return headers, rows, targets, exp


def build_lab7(folder):
    headers, rows, targets, _ = _lab7()
    wb = Workbook()
    ws = wb.active
    ws.title = "Applications"
    _write_rows(ws, headers, rows, "tblApps", {"Sourcing_Cost": '"S$"#,##0'}, {"Sector": 18, "Source": 19, "Recruiter": 15})
    ws2 = wb.create_sheet("Targets")
    _write_rows(ws2, ["Hire_Month", "Hire_Target"], targets, "tblTargets")
    wb.create_sheet("Dashboard")
    _about_sheet(wb, "Recruitment KPI Dashboard Data (Jan–Jun 2026)",
                 "Purpose: one row per application in H1 2026 with the furthest stage reached and the offer outcome. "
                 "Use it in Lab 7 to build a highly visual, interactive recruitment dashboard (KPI cards, PivotTables, "
                 "PivotCharts, slicers and a timeline) on the Dashboard sheet with Copilot, workbook rules and skills.",
                 [("Application_ID", "Text", "Unique key."), ("Applied_Date / Applied_Month", "Date / Text", "Application date and yyyy-mm label."),
                  ("Sector / Source / Recruiter", "Text", "Dashboard dimensions."),
                  ("Furthest_Stage", "Text", "Applied, Screened, Interviewed, Offered or Hired."),
                  ("Offer_Made", "Yes/No", "Whether an offer was extended."),
                  ("Offer_Accepted", "Yes/No/Pending", "Blank when no offer was made. Pending = decision after 2026-06-30."),
                  ("Hire_Date / Time_to_Hire_Days", "Date / Days", "Filled only for hires."),
                  ("Sourcing_Cost", "S$", "Channel cost attributed to this application."),
                  ("KPI definitions", "Rule", "Offer acceptance rate = accepted ÷ (accepted + declined); Cost per hire = total "
                   "Sourcing_Cost ÷ hires; Time to hire = average Time_to_Hire_Days of hires; hires are counted by Hire_Date month.")])
    _prompt_log_sheet(wb)
    return [_save(wb, folder, "BrightPath_Recruitment_Dashboard_Data.xlsx")]


# =========================================================================== LAB 8 (workforce dashboard skill)
DEPTS = ["Recruitment", "Sales & Client Services", "Payroll Services", "Finance", "Human Resources", "IT", "Operations"]


def _active_on(r, d):
    return r[5] <= d and (r[6] is None or r[6] > d)


def _lab8():
    rng = random.Random(3406)
    locs = ["Raffles Place HQ", "Jurong East Hub", "Tampines Hub", "Woodlands Hub"]
    dept_w = [30, 18, 14, 8, 7, 8, 15]
    level_w = {"Associate": 45, "Senior Associate": 28, "Manager": 18, "Senior Manager": 7, "Director": 2}
    level_sal = {"Associate": 3600, "Senior Associate": 5000, "Manager": 7400, "Senior Manager": 10200, "Director": 15500}
    rows = []
    for i in range(1, 421):
        d = rng.choices(DEPTS, dept_w)[0]
        loc = rng.choices(locs, [50, 18, 18, 14])[0]
        lvl = rng.choices(list(level_w), list(level_w.values()))[0]
        et = rng.choices(["Full-time", "Part-time", "Contract"], [80, 8, 12])[0]
        hire = _rand_date(rng, date(2019, 1, 7), date(2026, 6, 1))
        hazard = 0.26 if d in ("Recruitment", "Sales & Client Services") else 0.11
        exit_d, etype, status = None, "", "Active"
        cand = hire + timedelta(days=max(90, int(rng.expovariate(hazard) * 365)))
        if cand <= REPORT_DATE:
            exit_d = cand
            etype = rng.choices(["Voluntary", "Involuntary"], [82, 18])[0]
            status = "Exited"
        sal = int(round(level_sal[lvl] * rng.uniform(0.88, 1.15), -1))
        eng = round(min(5, max(1, rng.gauss(3.7 if status == "Active" else 3.0, 0.6))), 1)
        trn = max(0, int(rng.gauss(18, 8)))
        rows.append([f"EMP-{i:04d}", d, loc, lvl, et, hire, exit_d, etype, status, sal, eng, trn])
    rows.sort(key=lambda r: r[5])
    for i, r in enumerate(rows, 1):
        r[0] = f"EMP-{i:04d}"
    headers = ["Employee_ID", "Department", "Location", "Job_Level", "Employment_Type", "Hire_Date", "Exit_Date",
               "Exit_Type", "Status", "Monthly_Salary", "Engagement_Score", "Training_Hours_YTD"]
    p0, p1 = date(2025, 7, 1), REPORT_DATE
    active = [r for r in rows if r[8] == "Active"]
    start_hc = sum(1 for r in rows if _active_on(r, p0))
    end_hc = len(active)
    exits = [r for r in rows if r[6] and p0 <= r[6] <= p1]
    vol = [r for r in exits if r[7] == "Voluntary"]
    avg_hc = (start_hc + end_hc) / 2
    hires_12m = [r for r in rows if p0 <= r[5] <= p1]
    by_dept = {}
    for d in DEPTS:
        s0 = sum(1 for r in rows if r[1] == d and _active_on(r, p0))
        s1 = sum(1 for r in active if r[1] == d)
        ex = sum(1 for r in exits if r[1] == d)
        by_dept[d] = dict(hc=s1, exits=ex, rate=round(ex / ((s0 + s1) / 2) * 100, 1) if (s0 + s1) else 0)
    tenure = statistics.mean((REPORT_DATE - r[5]).days / 365.25 for r in active)
    new_joiners = []
    rng2 = random.Random(34061)
    for j in range(1, 9):
        d = rng2.choice(DEPTS[:3] + ["IT"])
        new_joiners.append([f"EMP-{len(rows) + j:04d}", d, rng2.choice(locs), "Associate", "Full-time",
                            date(2026, 7, rng2.choice([1, 6, 13, 20])), None, "", "Active",
                            int(round(3600 * rng2.uniform(0.9, 1.1), -1)), None, 0])
    exp = dict(rows=len(rows), active=end_hc, start_hc=start_hc, exits_12m=len(exits), vol_exits=len(vol),
               attrition=round(len(exits) / avg_hc * 100, 1), vol_attrition=round(len(vol) / avg_hc * 100, 1),
               hires_12m=len(hires_12m), by_dept=by_dept,
               top_attr_dept=max(by_dept, key=lambda k: by_dept[k]["rate"]), avg_tenure=round(tenure, 1),
               largest_dept=max(by_dept, key=lambda k: by_dept[k]["hc"]), new_joiners=len(new_joiners),
               avg_eng_active=round(statistics.mean(r[10] for r in active), 2),
               nj_by_dept={d: sum(1 for r in new_joiners if r[1] == d) for d in DEPTS if any(r[1] == d for r in new_joiners)},
               active_after=end_hc + len(new_joiners), rows_after=len(rows) + len(new_joiners))
    return headers, rows, new_joiners, exp


def build_lab8(folder):
    headers, rows, new_joiners, _ = _lab8()
    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"
    _write_rows(ws, headers, rows, "tblEmployees", {"Monthly_Salary": '"S$"#,##0', "Engagement_Score": "0.0"},
                {"Department": 22, "Location": 18, "Job_Level": 16})
    ws2 = wb.create_sheet("New_Joiners_Jul2026")
    _write_rows(ws2, headers, new_joiners, None, {"Monthly_Salary": '"S$"#,##0'}, {"Department": 22, "Location": 18})
    ws2["N1"] = "Do NOT use until the Lab 8 refresh test — these rows are appended to tblEmployees to test dashboard refresh."
    ws2["N1"].font = Font(bold=True, color="C00000")
    wb.create_sheet("HR_Dashboard")
    _about_sheet(wb, "Internal Workforce Snapshot (as at 30 Jun 2026)",
                 "Purpose: one row per BrightPath employee hired since January 2019 (Status Active or Exited). Use it in "
                 "Lab 8 to build an HR workforce dashboard with a custom Copilot skill. Salary and engagement data are "
                 "sensitive: share only aggregated views and never single out an individual.",
                 [("Employee_ID", "Text", "Pseudonymised employee key (no names are stored)."),
                  ("Department / Location / Job_Level / Employment_Type", "Text", "Dashboard dimensions."),
                  ("Hire_Date / Exit_Date", "Date", "Exit_Date blank for active employees."),
                  ("Exit_Type", "Text", "Voluntary or Involuntary; blank if active."), ("Status", "Text", "Active or Exited."),
                  ("Monthly_Salary", "S$", "Monthly base salary."), ("Engagement_Score", "1–5", "Latest pulse-survey score."),
                  ("Training_Hours_YTD", "Hours", "Training hours this calendar year."),
                  ("KPI definitions", "Rule", "12-month period = 2025-07-01 to 2026-06-30. Average headcount = (active on "
                   "2025-07-01 + active on 2026-06-30) ÷ 2. Attrition rate = exits in period ÷ average headcount.")])
    _prompt_log_sheet(wb)
    return [_save(wb, folder, "BrightPath_Workforce_Headcount.xlsx")]


# =========================================================================== registry
_BUILDERS = {1: build_lab1, 2: build_lab2, 3: build_lab3, 4: build_lab4,
             5: build_lab5, 6: build_lab6, 7: build_lab7, 8: build_lab8}

EXPECTED = {1: _lab1()[2], 2: _lab2()[2], 3: _lab3()[3], 4: _lab4()[3],
            5: _lab5()[3], 6: _lab6()[2], 7: _lab7()[3], 8: _lab8()[3]}


def build_lab_data(num, folder):
    return _BUILDERS[num](folder) if num in _BUILDERS else []


if __name__ == "__main__":
    import json
    print(json.dumps(EXPECTED, indent=1, default=str))
