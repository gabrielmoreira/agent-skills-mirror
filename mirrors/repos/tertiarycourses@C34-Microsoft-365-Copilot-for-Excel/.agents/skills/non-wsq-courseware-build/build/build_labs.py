#!/usr/bin/env python3
"""
Generate one folder per lab from the SAME single source (course_data.py +
data_domainN.py) that drives the PPT, LP and LG:

  labs/lab-NN-<slug>/
    Lab-NN-Instructions.pdf   the same instructions + sample prompts as a PDF
    sample-prompts.txt        copy-ready prompts
    data/<workbook>.xlsx      synthetic mock data (from mock_data.py)
    <extra files>             optional: e.g. skills/<name>/SKILL.md or a rules template (lab["extra_files"])

plus labs/README.md as the index. Lab files can never drift out of alignment
with the other artifacts.
"""
import os, re, sys, glob, importlib
from xml.sax.saxutils import escape

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import course_data as C
import mock_data
from lab_prompts import prompt_rows, reflow
from openpyxl import load_workbook
from openpyxl.styles import Alignment

from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (KeepTogether, ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle)


def find_repo(start):
    env = os.environ.get("COURSE_REPO")
    if env and os.path.isdir(env):
        return env
    d = start
    for _ in range(8):
        d = os.path.dirname(d)
        if os.path.isdir(os.path.join(d, "courseware")) and os.path.isdir(os.path.join(d, "labs")):
            return d
    return os.path.dirname(os.path.dirname(start))


REPO = find_repo(HERE)
LABS = os.path.join(REPO, "labs")
os.makedirs(LABS, exist_ok=True)

# Labs whose steps include prompts for a web chat (Copilot on the web) as well as for Excel
WEB_LABS = {6, 8}

# ---- collect labs from every data_domainN.py, in domain order
labs = []
for path in sorted(glob.glob(os.path.join(HERE, "data_domain[0-9]*.py"))):
    mod = importlib.import_module(os.path.basename(path)[:-3])
    for name in dir(mod):
        if re.fullmatch(r"DOMAIN\d+", name):
            labs.extend(getattr(mod, name))
labs.sort(key=lambda l: l["num"])
TOPICS = {t["num"]: t for t in C.TOPICS}


def slug(s):
    s = s.lower().replace("&", " and ")
    return re.sub(r"-+", "-", re.sub(r"[^a-z0-9]+", "-", s)).strip("-")


def folder_name(lab):
    # Folder named after the short lab theme (falls back to the full title).
    return f"lab-{lab['num']:02d}-{slug(lab.get('theme') or lab['title'])}"


def pdf_name(lab):
    return f"Lab-{lab['num']:02d}-Instructions.pdf"


def fence(cmd):
    return "```text\n" + cmd.strip() + "\n```"


# =========================================================================== PDF
def _font(name, candidates, fallback):
    for p in candidates:
        if os.path.exists(p):
            try:
                pdfmetrics.registerFont(TTFont(name, p))
                return name
            except Exception:
                pass
    return fallback


SANS = _font("C34Sans", ["/System/Library/Fonts/Supplemental/Arial.ttf", "C:/Windows/Fonts/arial.ttf",
                         "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"], "Helvetica")
SANS_B = _font("C34SansB", ["/System/Library/Fonts/Supplemental/Arial Bold.ttf", "C:/Windows/Fonts/arialbd.ttf",
                            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"], "Helvetica-Bold")
MONO = _font("C34Mono", ["/System/Library/Fonts/Supplemental/Courier New.ttf", "C:/Windows/Fonts/cour.ttf",
                         "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf"], "Courier")
if SANS != "Helvetica" and SANS_B != "Helvetica-Bold":
    from reportlab.lib.fonts import addMapping
    addMapping(SANS, 0, 0, SANS)
    addMapping(SANS, 1, 0, SANS_B)

BLUE = colors.HexColor("#1F6FEB")
NAVY = colors.HexColor("#0B3060")
INK = colors.HexColor("#111827")
GREY = colors.HexColor("#555B66")
LIGHT = colors.HexColor("#EEF4FF")
PROMPT_BG = colors.HexColor("#F3F6FA")
GREEN_BG = colors.HexColor("#E8F6EE")
GREEN = colors.HexColor("#1A7F4B")
AMBER_BG = colors.HexColor("#FFF6E5")
LINE = colors.HexColor("#D0D7E2")

ST = {
    "title": ParagraphStyle("t", fontName=SANS_B, fontSize=19, leading=23, textColor=INK, spaceAfter=4),
    "kicker": ParagraphStyle("k", fontName=SANS_B, fontSize=9.5, leading=12, textColor=BLUE, spaceAfter=2),
    "h2": ParagraphStyle("h2", fontName=SANS_B, fontSize=12.5, leading=16, textColor=NAVY, spaceBefore=10, spaceAfter=4),
    "h3": ParagraphStyle("h3", fontName=SANS_B, fontSize=10.5, leading=14, textColor=INK, spaceBefore=6, spaceAfter=2),
    "body": ParagraphStyle("b", fontName=SANS, fontSize=9.6, leading=13.4, textColor=INK, spaceAfter=4, alignment=TA_LEFT),
    "small": ParagraphStyle("s", fontName=SANS, fontSize=8.6, leading=11.5, textColor=INK),
    "smallb": ParagraphStyle("sb", fontName=SANS_B, fontSize=8.6, leading=11.5, textColor=INK),
    "code": ParagraphStyle("c", fontName=MONO, fontSize=8.3, leading=10.6, textColor=NAVY),
    "label": ParagraphStyle("l", fontName=SANS_B, fontSize=8, leading=10, textColor=BLUE),
}


def P(text, style="body"):
    return Paragraph(escape(text), ST[style])


def code_box(text, label=None):
    rows = []
    if label:
        rows.append([Paragraph(escape(label), ST["label"])])
    room = 170 * mm - 14  # column width minus left/right padding
    for line in text.strip("\n").split("\n"):
        lead = len(line) - len(line.lstrip(" "))
        style = ST["code"]
        # a formula has no spaces to wrap at, so a long one would split mid-token and break
        # when copied: shrink it to fit on one line instead
        width = pdfmetrics.stringWidth(line, MONO, style.fontSize)
        if line.lstrip().startswith("=") and width > room:
            size = max(6.0, style.fontSize * room / width * 0.95)  # margin: Paragraph wraps a little early
            style = ParagraphStyle("cf", parent=style, fontSize=size, leading=max(style.leading, size + 2))
        rows.append([Paragraph("&nbsp;" * lead + escape(line.lstrip(" ")) if line.strip() else "&nbsp;", style)])
    t = Table(rows, colWidths=[170 * mm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), PROMPT_BG),
        ("LINEBEFORE", (0, 0), (0, -1), 2.5, BLUE),
        ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 0.6), ("BOTTOMPADDING", (0, 0), (-1, -1), 0.6),
        ("TOPPADDING", (0, 0), (-1, 0), 5), ("BOTTOMPADDING", (0, -1), (-1, -1), 5),
    ]))
    return t


def callout(title, text, bg, edge):
    t = Table([[Paragraph(escape(title), ST["smallb"])], [Paragraph(escape(text), ST["small"])]], colWidths=[170 * mm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), bg), ("LINEBEFORE", (0, 0), (0, -1), 2.5, edge),
        ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return t


def grid(header, rows, widths):
    data = [[Paragraph(escape(h), ST["smallb"]) for h in header]]
    data += [[Paragraph(escape(str(c)), ST["small"]) for c in r] for r in rows]
    t = Table(data, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), LIGHT), ("GRID", (0, 0), (-1, -1), 0.4, LINE),
        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    return t


def bullets(items):
    return ListFlowable([ListItem(P(i), leftIndent=10, value="•") for i in items], bulletType="bullet", start="•",
                        leftIndent=12, bulletFontName=SANS, bulletFontSize=8)


def build_pdf(lab, topic, path):
    n = lab["num"]

    def on_page(canvas, doc):
        canvas.saveState()
        w, h = A4
        canvas.setFillColor(BLUE)
        canvas.rect(0, h - 9 * mm, w, 9 * mm, stroke=0, fill=1)
        canvas.setFillColor(colors.white)
        canvas.setFont(SANS_B, 8.5)
        canvas.drawString(20 * mm, h - 6 * mm, f"{C.TITLE}  ·  {C.COURSE_CODE}")
        canvas.drawRightString(w - 20 * mm, h - 6 * mm, f"Lab {n} of {len(labs)}")
        canvas.setFillColor(GREY)
        canvas.setFont(SANS, 7.5)
        canvas.drawString(20 * mm, 10 * mm, f"{C.ORG}  ·  {C.VERSION} ({C.VERSION_DATE})  ·  Fictional training data")
        canvas.drawRightString(w - 20 * mm, 10 * mm, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=17 * mm,
                            bottomMargin=17 * mm, title=f"Lab {n} — {lab['title']}", author=C.ORG,
                            subject=f"{C.TITLE} ({C.COURSE_CODE}) lab instructions and sample prompts")
    S = []
    S.append(Paragraph(escape(f"TOPIC {topic.get('code', '')} · {topic.get('title', '').upper()}"), ST["kicker"]))
    if lab.get("theme"):
        S.append(Paragraph(escape(f"LAB {n} THEME · {lab['theme'].upper()}"), ST["kicker"]))
    S.append(Paragraph(escape(f"Lab {n} — {lab['title']}"), ST["title"]))
    meta = [
        ("Maps to", lab["objective"]),
        ("Duration", f"{lab.get('duration', '—')} minutes"),
        ("Tools", lab["services"]),
        ("Mock data", ", ".join(lab.get("data_files", [])) or "—"),
    ]
    if lab.get("extra_files"):
        meta.append(("Lab files", "; ".join(f"{rel} ({desc})" for rel, _text, desc in lab["extra_files"])))
    meta += [
        ("Course", f"{C.TITLE} ({C.COURSE_CODE}) · {C.VERSION}"),
    ]
    rows = [[Paragraph(escape(k), ST["smallb"]), Paragraph(escape(v), ST["small"])] for k, v in meta]
    url = f"{C.PROMPTS_PAGE}#lab-{n}"
    rows.append([Paragraph("Prompts", ST["smallb"]),
                 Paragraph(f'<link href="{url}" color="#1F6FEB"><u>{escape(url)}</u></link>', ST["small"])])
    mt = Table(rows, colWidths=[26 * mm, 144 * mm])
    mt.setStyle(TableStyle([("BACKGROUND", (0, 0), (0, -1), LIGHT), ("GRID", (0, 0), (-1, -1), 0.4, LINE),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 3),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
    S += [Spacer(1, 4), mt]
    S += [Paragraph("Goal", ST["h2"]), P(lab.get("goal", lab["desc"]))]
    S += [Paragraph("What You Will Do", ST["h2"]), P(lab["desc"])]
    S += [Paragraph("What You Will Build", ST["h2"]), P(lab["build"])]
    if lab.get("prerequisites"):
        S += [Paragraph("Prerequisites", ST["h2"]), bullets(lab["prerequisites"])]
    S += [Spacer(1, 6), callout("Data note", C.LAB_NOTE, AMBER_BG, colors.HexColor("#E0A100"))]

    for i, (instr, cmd) in enumerate(lab["steps"], 1):
        cmd = reflow(cmd) if cmd.strip() else cmd
        block = [Paragraph(f"<b>Step {i}.</b> " + escape(instr), ST["body"])]
        if i == 1:
            block.insert(0, Paragraph("Steps", ST["h2"]))      # never strand the heading at a page foot
        if cmd.strip():
            block.append(code_box(cmd, "PROMPT / CHECK — copy and adapt"))
        block.append(Spacer(1, 5))
        if len(cmd) < 900:
            S.append(KeepTogether(block))
        else:  # long prompt box may split across pages; keep the heading and step text together
            lead = 2 if i == 1 else 1
            S += [KeepTogether(block[:lead])] + block[lead:]

    S.append(Paragraph("Sample Prompts", ST["h2"]))
    S.append(P(("Use the app labelled for each prompt (or adapt it). " if n in WEB_LABS else "Use these prompts in the Copilot pane (or adapt them). ")
               + "Name the table and columns, state the output, and "
               "always verify the result. The same prompts are in sample-prompts.txt in this lab folder."))
    url = f"{C.PROMPTS_PAGE}#lab-{n}"
    S.append(Paragraph(f'<b>Copy-ready prompts online:</b> <link href="{url}" color="#1F6FEB"><u>{escape(url)}</u></link>',
                       ST["body"]))
    sp = lab.get("sample_prompts", [])
    S.append(grid(["#", "Purpose", "Prompt"], [(str(i), lbl, pr) for i, (lbl, pr) in enumerate(sp, 1)],
                  [8 * mm, 28 * mm, 134 * mm]))

    # Keep each short closing section with its heading so no heading is orphaned.
    if lab.get("checks"):
        S.append(KeepTogether([Paragraph("Test It", ST["h2"]),
                               grid(["Check", "Expected result"], lab["checks"], [62 * mm, 108 * mm])]))
    else:
        S.append(KeepTogether([Paragraph("Test It", ST["h2"]), callout("Expected results", lab["test"], GREEN_BG, GREEN)]))
    if lab.get("checkpoint"):
        S.append(KeepTogether([Paragraph("Checkpoint", ST["h2"]), P(lab["checkpoint"])]))
    if lab.get("troubleshooting"):
        S.append(KeepTogether([Paragraph("Troubleshooting", ST["h2"]),
                               grid(["If this happens", "Fix"], lab["troubleshooting"], [58 * mm, 112 * mm])]))
    if lab.get("challenge"):
        S.append(KeepTogether([Paragraph("Challenge", ST["h2"]), P(lab["challenge"])]))
    if lab.get("reflection"):
        S.append(KeepTogether([Paragraph("Reflection", ST["h2"]), P(lab["reflection"])]))
    doc.build(S, onFirstPage=on_page, onLaterPages=on_page)


def build_prompts_txt(lab):
    out = [f"{C.TITLE} ({C.COURSE_CODE}) — Lab {lab['num']}: {lab['title']}",
           ("Sample prompts: use the app labelled in each step. Copy one at a time, adapt it, and verify every result."
            if lab["num"] in WEB_LABS else "Sample prompts for Copilot in Excel. Copy one at a time, adapt it, and verify every result."),
           f"Prompts web page (one-click Copy): {C.PROMPTS_PAGE}#lab-{lab['num']}", ""]
    for i, (lbl, pr) in enumerate(lab.get("sample_prompts", []), 1):
        out += [f"[{i}] {lbl}", pr, ""]
    out += ["-" * 60, "Step prompts used in this lab", "-" * 60, ""]
    for i, (_instr, cmd) in enumerate(lab["steps"], 1):
        if cmd.strip():
            out += [f"[Step {i}]", reflow(cmd).strip(), ""]
    return "\n".join(out)


# =========================================================================== Prompt_Log
def fill_prompt_log(path, lab):
    """Pre-fill the workbook's Prompt_Log with every prompt of the lab, in order (Step, Mode, Prompt).
    The learner asks Copilot to fill in What Copilot did, How I verified and Decision."""
    wb = load_workbook(path)
    if "Prompt_Log" not in wb.sheetnames:
        return
    ws = wb["Prompt_Log"]
    if lab["num"] in WEB_LABS:
        ws.cell(row=1, column=2, value="App / mode (Excel or Copilot web)")
    rows = [(f"{r['step']} · {r['letter']}", r["mode"], r["text"]) for r in prompt_rows(lab)]
    rows += [(f"Extra {i}", "Any", pr) for i, (_lbl, pr) in enumerate(lab.get("sample_prompts", []), 1)]
    top = Alignment(wrap_text=True, vertical="top")
    for r, vals in enumerate(rows, 2):
        for c, v in enumerate(vals, 1):
            ws.cell(row=r, column=c, value=v).alignment = top
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 80
    wb.save(path)


# =========================================================================== build
written = []
names = [folder_name(l) for l in labs]
for idx, lab in enumerate(labs):
    topic = TOPICS.get(lab["topic"], {})
    fname = names[idx]
    folder = os.path.join(LABS, fname)
    os.makedirs(folder, exist_ok=True)
    mock_data.build_lab_data(lab["num"], folder)
    for rel, text, _desc in lab.get("extra_files", []):
        p = os.path.join(folder, rel)
        os.makedirs(os.path.dirname(p), exist_ok=True)
        with open(p, "w", encoding="utf-8") as f:
            f.write(text)
    for rel in lab.get("data_files", []):
        assert os.path.exists(os.path.join(folder, rel)), f"missing mock data {rel} for lab {lab['num']}"
        fill_prompt_log(os.path.join(folder, rel), lab)
    with open(os.path.join(folder, "sample-prompts.txt"), "w", encoding="utf-8") as f:
        f.write(build_prompts_txt(lab))
    build_pdf(lab, topic, os.path.join(folder, pdf_name(lab)))
    written.append((lab["num"], fname, lab))

# ---- labs/README.md index
R = [f"# {C.TITLE} — Hands-On Labs\n",
     f"{len(labs)} labs across {len(C.TOPICS)} topics · {C.DAYS} day · "
     f"{C.DAYS * getattr(C, 'INSTRUCTIONAL_MINUTES', C.DAY_MINUTES) / 60:g} instructional hours\n",
     "Every lab has its own folder containing an **instruction PDF with sample prompts**, "
     "a copy-ready `sample-prompts.txt`, and a `data/` folder with the **mock HR / recruitment workbook**. "
     "All data belongs to *BrightPath Talent Solutions*, a fictional Singapore recruitment and staffing firm.\n",
     f"**Prompts web page:** every lab's prompts with one-click Copy buttons — [{C.PROMPTS_PAGE}]({C.PROMPTS_PAGE})\n",
     "Work through the labs in order. Each lab opens its own workbook, so you can rejoin at any lab.\n"]
for t in C.TOPICS:
    R.append(f"\n## Topic {t['num']} — {t['title']}\n")
    R.append("| # | Theme | Lab | Instructions | Mock data | You Build |")
    R.append("|---|-------|-----|--------------|-----------|-----------|")
    for n, fname, lab in written:
        if lab["topic"] == t["num"]:
            data = "<br>".join([f"[{os.path.basename(d)}]({fname}/{d})" for d in lab.get("data_files", [])]
                               + [f"[{rel}]({fname}/{rel})" for rel, _t, _d in lab.get("extra_files", [])])
            R.append(f"| {n} | {lab.get('theme', '')} | [{lab['title']}]({fname}/{pdf_name(lab)}) | [PDF]({fname}/{pdf_name(lab)}) · "
                     f"[prompts]({fname}/sample-prompts.txt) | {data} | {lab['build']} |")
R.append(f"\n---\n\n> {C.LAB_NOTE}\n")
R.append(f"\n_{C.ORG} · {C.COURSE_CODE} · {C.VERSION} ({C.VERSION_DATE})_\n")
with open(os.path.join(LABS, "README.md"), "w", encoding="utf-8") as f:
    f.write("\n".join(R))

print(f"Saved {len(labs)} lab folders (PDF, prompts, data) + README.md to {LABS}")
